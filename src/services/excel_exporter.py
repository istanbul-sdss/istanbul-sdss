"""
src/services/excel_exporter.py

Production-level çok sayfalı Excel raporu.
Şu sayfaları üretir:
  📊 Özet          — kategori bazında kayıt sayısı + güven dağılımı
  🗺️ Mahalle        — mahalle × kategori pivot
  📍 <Kategori>     — her kategori için detay sayfası
  🔍 Veri Kalitesi  — doluluk oranları, güven dağılımı
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.logger import get_logger

log = get_logger(__name__)

# ── STİL YARDIMCILARI ────────────────────────────────────────────────────────
_thin   = Side(style="thin", color="D0D0D0")
_BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

GROUP_BG = {
    "Sağlık": "C0392B", "Eğitim": "D68910", "Binalar": "2471A3",
    "Yeşil Alan": "1E8449", "Ulaşım": "1A5276", "Ticaret": "784212",
    "Altyapı": "616A6B", "Kültür": "6C3483", "Kamu": "2C3E50",
}



def _hdr(ws, row, col, val, bg="1B4F72", sz=11):
    c = ws.cell(row, col, val)
    c.fill   = PatternFill("solid", fgColor=bg)
    c.font   = Font(bold=True, color="FFFFFF", size=sz)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = _BORDER
    return c


def _dat(ws, row, col, val, alt=False, bold=False, color=None):
    v = "" if (isinstance(val, float) and val != val) else val
    c = ws.cell(row, col, v)
    bg = color if color else ("EBF5FB" if alt else "FFFFFF")
    c.fill   = PatternFill("solid", fgColor=bg)
    c.alignment = Alignment(vertical="center")
    c.border = _BORDER
    c.font   = Font(size=10, bold=bold)
    return c


def _auto_width(ws, df: pd.DataFrame, start_col: int = 1):
    for ci, col in enumerate(df.columns, start_col):
        q = df[col].astype(str).str.len().quantile(0.9) if len(df) > 0 else 0
        q_int = int(q) if q == q else 0   # NaN kontrolü (NaN != NaN)
        max_w = max(len(str(col)), q_int) + 3
        ws.column_dimensions[get_column_letter(ci)].width = min(max_w, 45)


# ══════════════════════════════════════════════════════════════════════════════
class ExcelExporter:

    def __init__(self, output_dir: str | None = None):
        from src.config.settings import OUTPUT_DIR
        self.out = Path(output_dir) if output_dir else OUTPUT_DIR
        self.out.mkdir(parents=True, exist_ok=True)

    def export(
        self,
        results: dict,
        ilce:    str,
    ) -> str:
        """
        Tam rapor üretir.

        `results` artık pipeline'dan gelen rule_code-keyed dict; her değer
        en az `df`, `label_tr`, `category_group` taşır:

            results = {
                "pharmacy": {"df": ..., "label_tr": "Eczane", ...},
                "building_residential": {"df": ..., "label_tr": "Konut - Genel", ...},
            }

        P1.1 düzeltmesi: rule_code teknik kimliktir; sheet adı, özet
        "Kategori" sütunu, başlıklar gibi KULLANICI GÖRÜR yüzeyler artık
        `label_tr` kullanır. Çakışan label_tr'ler "Karakol", "Karakol 2"
        gibi otomatik benzersizleştirilir.
        """
        wb = Workbook()
        wb.remove(wb.active)

        # Boş olmayan kategoriler — anahtar hâlâ rule_code (debug/log için)
        nonempty = {k: v for k, v in results.items() if not v["df"].empty}

        # rule_code → display_name (label_tr) eşlemesi, çakışmalar
        # benzersizleştirilmiş halde. Kullanıcıya gösterilen her yerde bu kullanılır.
        display_names = self._build_display_names(nonempty)

        # display_name-keyed dict'ler (sheet adı = display_name); her _ozet/_kategori
        # bu hali tüketir.
        dataframes = {display_names[k]: v["df"]              for k, v in nonempty.items()}
        summaries  = {display_names[k]: v.get("summary", {}) for k, v in nonempty.items()}

        # ── Sayfalar ──────────────────────────────────────────────────────────
        self._ozet(wb, dataframes, summaries, ilce)

        all_df = pd.concat(dataframes.values(), ignore_index=True) if dataframes else pd.DataFrame()

        if not all_df.empty and "Mahalle" in all_df.columns:
            self._mahalle_pivot(wb, all_df, dataframes)

        for display_name, df in dataframes.items():
            self._kategori(wb, df, display_name, ilce)

        if not all_df.empty:
            self._veri_kalitesi(wb, dataframes, ilce)

        # ── Kaydet ────────────────────────────────────────────────────────────
        ts   = datetime.now().strftime("%Y%m%d_%H%M")
        path = self.out / f"{ilce}_OSM_{ts}.xlsx"
        wb.save(path)
        log.info(f"Excel kaydedildi: {path}")
        return str(path)

    @staticmethod
    def _build_display_names(nonempty: dict) -> dict[str, str]:
        """
        rule_code → benzersiz label_tr eşlemesi.

        Aynı label_tr birden fazla rule'da geçerse (ör. "Karakol" hem
        building_police hem infrastructure_police için) sırayla " 2", " 3"
        suffix eklenir. Pipeline'ın `res["label_tr"]` sözleşmesini onurlandırır.
        """
        used: dict[str, int] = {}
        out: dict[str, str] = {}
        for rule_code, res in nonempty.items():
            base = (res.get("label_tr") or rule_code).strip() or rule_code
            n = used.get(base, 0) + 1
            used[base] = n
            out[rule_code] = base if n == 1 else f"{base} {n}"
        return out

    # ── ÖZET ──────────────────────────────────────────────────────────────────
    def _ozet(self, wb, dataframes, summaries, ilce):
        ws = wb.create_sheet("📊 Özet")
        ws.sheet_view.showGridLines = False

        # Başlık
        n_col = 9
        ws.merge_cells(f"A1:{get_column_letter(n_col)}1")
        c = ws.cell(1, 1, f"İstanbul — {ilce} İlçesi Mekânsal Veri Raporu")
        c.font = Font(bold=True, size=16, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1B4F72")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 38

        ws.merge_cells(f"A2:{get_column_letter(n_col)}2")
        c2 = ws.cell(2, 1, f"Oluşturulma: {datetime.now().strftime('%d.%m.%Y %H:%M')}  |  Kaynak: OpenStreetMap / Overpass API  |  © OSM Katkıcıları")
        c2.font = Font(italic=True, size=9, color="555555")
        c2.alignment = Alignment(horizontal="center")
        ws.row_dimensions[2].height = 18
        ws.row_dimensions[3].height = 6

        hdrs = ["#", "Kategori", "Kayıt", "Mahalle",
                "Koordinatlı", "Alanlı", "Yüksek Güven", "Orta Güven", "Durum"]
        for ci, h in enumerate(hdrs, 1):
            _hdr(ws, 4, ci, h, "2E86C1")
        ws.row_dimensions[4].height = 26

        toplam = 0
        for sira, (kat, df) in enumerate(dataframes.items(), 1):
            row = 4 + sira
            alt = sira % 2 == 0
            n   = len(df)
            toplam += n
            mah   = df["Mahalle"].nunique() if "Mahalle" in df.columns else 0
            koord = df["Enlem"].notna().sum() if "Enlem" in df.columns else 0
            alan  = df["Alan (m²)"].notna().sum() if "Alan (m²)" in df.columns else 0
            yuk   = (df["Güven"] == "Yüksek").sum() if "Güven" in df.columns else 0
            ort   = (df["Güven"] == "Orta").sum()   if "Güven" in df.columns else 0
            for ci, v in enumerate([sira, kat, n, mah, koord, alan, yuk, ort,
                                     "✅ Tamam" if n > 0 else "⚠️ Veri Yok"], 1):
                _dat(ws, row, ci, v, alt)
            ws.row_dimensions[row].height = 20

        tr = 4 + len(dataframes) + 1
        ws.merge_cells(f"A{tr}:B{tr}")
        for ci in range(1, n_col+1):
            c = ws.cell(tr, ci)
            c.fill = PatternFill("solid", fgColor="1B4F72")
            c.font = Font(bold=True, color="FFFFFF", size=11)
            c.border = _BORDER
            c.alignment = Alignment(horizontal="center")
        ws.cell(tr, 1, "TOPLAM")
        ws.cell(tr, 3, toplam)

        for col, w in zip("ABCDEFGHI", [5, 30, 10, 12, 14, 12, 14, 12, 14]):
            ws.column_dimensions[col].width = w

        if len(dataframes) > 1:
            chart = BarChart()
            chart.type   = "col"
            chart.title  = f"{ilce} — Kategori Dağılımı"
            chart.y_axis.title = "Kayıt Sayısı"
            chart.style  = 10
            chart.width  = 22
            chart.height = 14
            data_ref = Reference(ws, min_col=3, max_col=3, min_row=4, max_row=4+len(dataframes))
            cat_ref  = Reference(ws, min_col=2, max_col=2, min_row=5, max_row=4+len(dataframes))
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(cat_ref)
            ws.add_chart(chart, f"A{tr+3}")

    # ── MAHALLE PİVOT ─────────────────────────────────────────────────────────
    def _mahalle_pivot(self, wb, all_df, dataframes):
        ws = wb.create_sheet("🗺️ Mahalle Özeti")
        ws.sheet_view.showGridLines = False

        kat_col = "Kategori (TR)" if "Kategori (TR)" in all_df.columns else None
        if kat_col:
            pivot = all_df.groupby(["Mahalle", kat_col]).size().unstack(fill_value=0)
        else:
            pivot = all_df.groupby("Mahalle").size().to_frame("Toplam")

        pivot["TOPLAM"] = pivot.sum(axis=1)
        pivot = pivot.sort_values("TOPLAM", ascending=False).reset_index()

        hdrs = list(pivot.columns)
        for ci, h in enumerate(hdrs, 1):
            _hdr(ws, 1, ci, h, "1E8449")
        ws.row_dimensions[1].height = 28

        for ri, row_data in pivot.iterrows():
            alt = ri % 2 == 0
            for ci, h in enumerate(hdrs, 1):
                v    = row_data.get(h, 0)
                bold = (h == "TOPLAM")
                _dat(ws, ri+2, ci, v, alt, bold=bold)
            ws.row_dimensions[ri+2].height = 18

        ws.column_dimensions["A"].width = 28
        for ci in range(2, len(hdrs)+1):
            ws.column_dimensions[get_column_letter(ci)].width = 16

    # ── KATEGORİ DETAY ────────────────────────────────────────────────────────
    def _kategori(self, wb, df: pd.DataFrame, kat_adi: str, ilce: str):
        # Excel sayfa adında geçersiz karakterleri temizle
        # openpyxl: / \ ? * [ ] : karakterleri yasak, max 31 karakter
        temiz_adi = (kat_adi
                     .replace("/", "-")
                     .replace("\\", "-")
                     .replace("?", "")
                     .replace("*", "")
                     .replace("[", "(")
                     .replace("]", ")")
                     .replace(":", "-")
                     .strip())
        sheet_title = f"📍 {temiz_adi}"[:31]

        ws = wb.create_sheet(sheet_title)
        ws.sheet_view.showGridLines = False

        bg = next((c for k, c in GROUP_BG.items() if k in kat_adi), "1B4F72")
        n  = len(df.columns)

        ws.merge_cells(f"A1:{get_column_letter(n)}1")
        c = ws.cell(1, 1, f"{ilce} — {kat_adi} ({len(df):,} kayıt)")
        c.font = Font(bold=True, size=13, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=bg)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        for ci, col in enumerate(df.columns, 1):
            _hdr(ws, 2, ci, col, bg)
        ws.row_dimensions[2].height = 22

        for ri, (_, row) in enumerate(df.iterrows(), 3):
            alt = ri % 2 == 0
            for ci, val in enumerate(row, 1):
                _dat(ws, ri, ci, val, alt)
            ws.row_dimensions[ri].height = 16

        ws.auto_filter.ref = f"A2:{get_column_letter(n)}{len(df)+2}"
        ws.freeze_panes = "A3"
        _auto_width(ws, df)

    # ── VERİ KALİTESİ ─────────────────────────────────────────────────────────
    def _veri_kalitesi(self, wb, dataframes: dict, ilce: str):
        ws = wb.create_sheet("🔍 Veri Kalitesi")
        ws.sheet_view.showGridLines = False

        ws.merge_cells("A1:G1")
        c = ws.cell(1, 1, f"{ilce} — Veri Kalite Raporu")
        c.font = Font(bold=True, size=13, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2C3E50")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        hdrs = ["Kategori", "Kayıt", "Mahalle Dolu%", "Koordinat%",
                "Alan%", "Yüksek Güven%", "Ort. Doluluk%"]
        for ci, h in enumerate(hdrs, 1):
            _hdr(ws, 2, ci, h, "34495E")

        for ri, (kat, df) in enumerate(dataframes.items(), 3):
            n = len(df)
            if n == 0:
                continue

            # B023: loop değişkeni df'i default-arg ile bağla.
            def pct(col, _df=df):
                if col not in _df.columns:
                    return 0
                return round(_df[col].notna().mean() * 100, 1)

            mah_pct   = pct("Mahalle")
            koord_pct = pct("Enlem")
            alan_pct  = pct("Alan (m²)")
            gvn_pct   = round((df["Güven"] == "Yüksek").mean() * 100, 1) \
                        if "Güven" in df.columns else 0

            tum_pct = round(
                (mah_pct + koord_pct + alan_pct) / 3, 1
            )

            alt = ri % 2 == 0
            for ci, v in enumerate([kat, n, mah_pct, koord_pct, alan_pct, gvn_pct, tum_pct], 1):
                cell = _dat(ws, ri, ci, v, alt)
                if ci > 2 and isinstance(v, (int, float)):
                    color = "D5F5E3" if v >= 80 else "FDEBD0" if v >= 50 else "FADBD8"
                    cell.fill = PatternFill("solid", fgColor=color)

        for col, w in zip("ABCDEFG", [30, 10, 14, 14, 12, 16, 14]):
            ws.column_dimensions[col].width = w
