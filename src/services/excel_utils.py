"""
src/services/excel_utils.py

Stilize Excel workbook üretiminde tekrarlanan boilerplate'i tek noktaya alır.

Üç çağıran:
  • pages/1_Data_Extraction.py  → çok-sayfalı extraction raporu
  • pages/4_Name_Lookup.py      → tek-sayfalı lookup raporu
  • Optimization_Tool.py        → 9-sayfalı optimization raporu

Sayfa içerikleri (DataFrame'ler) çağıranlarda kalır; bu modül yalnızca
"openpyxl üzerinden tutarlı stil + kolon genişliği + freeze + autofilter"
kısmını ortaklaştırır.

Why: Stil bloğu üç dosyada %95 aynıydı. Header rengi/font değişikliği
üç yerde unutulup tutarsız çıktı üretiyordu.
"""

from __future__ import annotations

from collections.abc import Callable, Iterable
from io import BytesIO

import pandas as pd

# Tek noktada brand stili.
HEADER_FONT_KW = dict(bold=True, color="FFFFFF", size=11)
HEADER_FILL_HEX = "0F2A44"
TOTAL_FILL_HEX  = "DBEAFE"

DEFAULT_MAX_COL_WIDTH = 45
DEFAULT_MIN_COL_WIDTH = 10
HEADER_ROW_HEIGHT     = 24
COL_WIDTH_SCAN_ROWS   = 500   # Performans: ilk 500 satıra göre genişlik kestir.


def style_workbook(
    writer: pd.ExcelWriter,
    *,
    total_row_sheets: Iterable[str] = (),
    skip_autofilter_sheets: Iterable[str] = (),
    highlight_last_column_sheets: Iterable[str] = (),
) -> None:
    """
    Aktif `pd.ExcelWriter` (openpyxl engine) için kapanmadan ÖNCE çağrılır.

    Parametreler
    ----------
    writer
        Açık ExcelWriter — `with pd.ExcelWriter(...) as writer:` içinde.
    total_row_sheets
        Bu isimdeki sayfalarda 1. kolonu "TOTAL" olan satırlar
        kalın+mavi vurgulanır (Summary / Pivot için).
    skip_autofilter_sheets
        Autofilter eklenmesi istenmeyen sayfalar (genellikle Pivot tabloları —
        çoklu satır başlığı ile autofilter çakışıyor).
    highlight_last_column_sheets
        Son kolonu (TOTAL kolonu) vurgulanacak sayfalar.
    """
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(**HEADER_FONT_KW)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL_HEX)
    total_font  = Font(bold=True)
    total_fill  = PatternFill("solid", fgColor=TOTAL_FILL_HEX)
    header_align = Alignment(horizontal="center", vertical="center")

    total_set     = set(total_row_sheets)
    no_filter_set = set(skip_autofilter_sheets)
    last_col_set  = set(highlight_last_column_sheets)

    for sheet_name in writer.sheets:
        ws = writer.sheets[sheet_name]

        # 1. Header satırı stili
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
        ws.row_dimensions[1].height = HEADER_ROW_HEIGHT
        ws.freeze_panes = "A2"

        # 2. Kolon genişliği — başlık + ilk 500 satıra göre kestirme
        for col_idx in range(1, ws.max_column + 1):
            letter = get_column_letter(col_idx)
            max_len = DEFAULT_MIN_COL_WIDTH
            for row_idx in range(1, min(ws.max_row + 1, COL_WIDTH_SCAN_ROWS + 1)):
                v = ws.cell(row=row_idx, column=col_idx).value
                if v is not None:
                    max_len = max(max_len, min(len(str(v)) + 2, DEFAULT_MAX_COL_WIDTH))
            ws.column_dimensions[letter].width = max_len

        # 3. TOTAL satırı vurgusu
        if sheet_name in total_set:
            for row_idx in range(2, ws.max_row + 1):
                first = ws.cell(row=row_idx, column=1).value
                if str(first).strip().upper() == "TOTAL":
                    for col_idx in range(1, ws.max_column + 1):
                        c = ws.cell(row=row_idx, column=col_idx)
                        c.font = total_font
                        c.fill = total_fill

        # 4. Son kolon vurgusu (Pivot TOTAL kolonu vb.)
        if sheet_name in last_col_set and ws.max_column > 0:
            last = ws.max_column
            for row_idx in range(2, ws.max_row + 1):
                c = ws.cell(row=row_idx, column=last)
                c.font = total_font
                c.fill = total_fill

        # 5. Autofilter — tek satırlı header'ı olan sayfalarda anlamlı
        if sheet_name not in no_filter_set and ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions


def build_styled_workbook(
    write_sheets: Callable[[pd.ExcelWriter], None],
    *,
    total_row_sheets: Iterable[str] = (),
    skip_autofilter_sheets: Iterable[str] = (),
    highlight_last_column_sheets: Iterable[str] = (),
) -> bytes:
    """
    Tek seferlik kullanım için kısa yol: callback sayfaları yazar, helper
    BytesIO kurar + stil uygular + bytes döner.

    Örnek:
        def _write(w):
            df1.to_excel(w, sheet_name="Summary", index=False)
            df2.to_excel(w, sheet_name="Detail", index=False)

        bytes_out = build_styled_workbook(_write, total_row_sheets=["Summary"])
    """
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        write_sheets(writer)
        style_workbook(
            writer,
            total_row_sheets=total_row_sheets,
            skip_autofilter_sheets=skip_autofilter_sheets,
            highlight_last_column_sheets=highlight_last_column_sheets,
        )
    return buf.getvalue()
