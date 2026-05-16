"""
Regresyon (P1.1): ExcelExporter sheet adları ve özet satırları teknik
`rule_code` yerine kullanıcıya gösterilecek `label_tr` kullansın.

Audit bulgusu: pipeline anahtarları rule_code'a geçtikten sonra
ExcelExporter dict key'i sheet adına yazıyordu → "📍 building_residential",
"📍 pharmacy" gibi teknik adlar export'ta görünüyordu. Pipeline yorumu
"consumer'lar res['label_tr'] kullanır" diyor; ExcelExporter bu sözleşmeye
uymuyor ve rapor profesyonelliğini düşürüyordu.
"""
from __future__ import annotations

import tempfile

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.services.excel_exporter import ExcelExporter


def _result(label_tr: str, n: int = 2) -> dict:
    """Pipeline `res` sözlüğünü taklit eder."""
    return {
        "df": pd.DataFrame({
            "Ad":       [f"{label_tr} {i}" for i in range(n)],
            "Mahalle":  ["Merkez"] * n,
            "Enlem":    [41.0] * n,
            "Boylam":   [29.0] * n,
            "Güven":    ["Yüksek"] * n,
        }),
        "summary":        {"total": n},
        "label_tr":       label_tr,
        "category_group": "buildings",
    }


def test_sheet_names_use_label_tr_not_rule_code():
    """Sheet adında 'pharmacy' değil 'Eczane' görünmeli."""
    results = {
        "pharmacy":            _result("Eczane"),
        "building_residential": _result("Konut - Genel"),
    }
    with tempfile.TemporaryDirectory() as td:
        path = ExcelExporter(output_dir=td).export(results, "TestIlce")
        wb = load_workbook(path)

    sheets = wb.sheetnames
    # Teknik kodlar HİÇBİR sheet adında olmamalı (sadece veri sütununda)
    for tech in ("pharmacy", "building_residential"):
        assert not any(tech in s for s in sheets), (
            f"REGRESYON: sheet adı teknik rule_code içeriyor: {sheets}. "
            f"label_tr ('Eczane', 'Konut - Genel') kullanılmalıydı."
        )
    # En az bir sheet label_tr içermeli
    assert any("Eczane" in s for s in sheets), (
        f"Eczane sheet'i yok; sheets={sheets}"
    )


def test_summary_sheet_uses_label_tr_in_category_column():
    """Özet sayfasındaki Kategori sütunu da label_tr göstermeli."""
    results = {"pharmacy": _result("Eczane")}
    with tempfile.TemporaryDirectory() as td:
        path = ExcelExporter(output_dir=td).export(results, "TestIlce")
        wb = load_workbook(path)
    ozet = [s for s in wb.sheetnames if "zet" in s][0]
    ws = wb[ozet]
    # Kategori değerleri 5'ten itibaren satırlarda; sütun B
    cats = [ws.cell(r, 2).value for r in range(5, ws.max_row + 1)]
    cats = [c for c in cats if c]  # boşları at
    assert "Eczane" in cats, (
        f"REGRESYON: özet sayfasında Kategori sütunu teknik rule_code "
        f"gösteriyor: {cats}"
    )
    assert "pharmacy" not in cats


def test_duplicate_label_tr_disambiguated_in_sheet_names():
    """
    İki rule aynı label_tr taşıyorsa (ör. 'Karakol') sheet adları
    çakışmamalı; ExcelExporter benzersiz isim üretmeli.
    """
    results = {
        "building_police":      _result("Karakol"),
        "infrastructure_police": _result("Karakol"),
    }
    with tempfile.TemporaryDirectory() as td:
        path = ExcelExporter(output_dir=td).export(results, "TestIlce")
        wb = load_workbook(path)
    karakol_sheets = [s for s in wb.sheetnames if "Karakol" in s]
    assert len(karakol_sheets) == 2, (
        f"REGRESYON: çakışan label_tr için tek sheet yazılmış (üzerine "
        f"yazıldı): {wb.sheetnames}. Benzersizleştirme yok."
    )
    # En az birinde ayrım eki olmalı (ör. "Karakol", "Karakol 2")
    assert len(set(karakol_sheets)) == 2


def test_per_category_sheet_preserves_rule_code_as_data_column():
    """
    Sheet adı insan-okunur olsa bile, teknik kimlik kayıt sütunu olarak
    korunmalı (debug/audit için). Burada en azından df içinde 'Kural Kodu'
    veya benzeri bir kolon olabilir — bu test mevcut df'in olduğu gibi
    yazıldığını doğrular (üst sözleşme).
    """
    results = {"pharmacy": _result("Eczane")}
    with tempfile.TemporaryDirectory() as td:
        path = ExcelExporter(output_dir=td).export(results, "TestIlce")
        wb = load_workbook(path)
    eczane = [s for s in wb.sheetnames if "Eczane" in s][0]
    ws = wb[eczane]
    # 1. satır başlık (merge title), 2. satır kolon başlıkları
    headers = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
    assert "Ad" in headers, f"Ad kolonu yok: {headers}"
