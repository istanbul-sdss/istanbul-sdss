"""
tests/test_excel_styling.py

`src/services/excel_utils.py` regresyon testi.

Faz 2 konsolidasyonunda 3 farklı dosyadaki Excel styling pattern'i tek bir
helper'a indirgendi (`style_workbook` + `build_styled_workbook`). Bu testler
helper'ın sözleşmesini sabitler ki ileride brand stilini değiştirirken
(font/renk/genişlik) regresyon yakalansın.
"""
from __future__ import annotations

from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.services.excel_utils import (
    DEFAULT_MAX_COL_WIDTH,
    HEADER_FILL_HEX,
    HEADER_ROW_HEIGHT,
    TOTAL_FILL_HEX,
    build_styled_workbook,
    style_workbook,
)


def _load(bytes_out: bytes):
    return load_workbook(BytesIO(bytes_out))


# ─────────────────────────────────────────────────────────────────────────────
# build_styled_workbook — happy path
# ─────────────────────────────────────────────────────────────────────────────
def test_build_styled_workbook_returns_bytes():
    """Helper bytes dönmeli ve openpyxl ile açılabilmeli."""
    out = build_styled_workbook(
        lambda w: pd.DataFrame({"A": [1, 2], "B": [3, 4]}).to_excel(
            w, sheet_name="X", index=False
        )
    )
    assert isinstance(out, bytes)
    assert len(out) > 0
    wb = _load(out)
    assert "X" in wb.sheetnames


def test_header_styled_with_brand_colors():
    """Header satırı koyu mavi fill + beyaz bold font + ortalı olmalı."""
    out = build_styled_workbook(
        lambda w: pd.DataFrame({"Col1": [1], "Col2": [2]}).to_excel(
            w, sheet_name="X", index=False
        )
    )
    ws = _load(out)["X"]

    a1, b1 = ws["A1"], ws["B1"]
    # Fill — openpyxl 2 hex char alpha prefix ekleyebiliyor; suffix kontrol
    assert a1.fill.fgColor.rgb.upper().endswith(HEADER_FILL_HEX.upper())
    assert b1.fill.fgColor.rgb.upper().endswith(HEADER_FILL_HEX.upper())
    assert a1.font.bold is True
    assert a1.font.color.rgb.upper().endswith("FFFFFF")
    assert a1.alignment.horizontal == "center"
    assert a1.alignment.vertical == "center"
    assert ws.row_dimensions[1].height == HEADER_ROW_HEIGHT


def test_freeze_panes_set_to_a2():
    out = build_styled_workbook(
        lambda w: pd.DataFrame({"A": [1]}).to_excel(w, sheet_name="X", index=False)
    )
    ws = _load(out)["X"]
    assert ws.freeze_panes == "A2"


def test_autofilter_added_when_rows_present():
    """≥1 veri satırı varsa autofilter set edilmeli."""
    out = build_styled_workbook(
        lambda w: pd.DataFrame({"A": [1, 2], "B": [3, 4]}).to_excel(
            w, sheet_name="X", index=False
        )
    )
    ws = _load(out)["X"]
    assert ws.auto_filter.ref is not None


def test_autofilter_skipped_for_listed_sheets():
    """`skip_autofilter_sheets` → autofilter pas geçilmeli (Pivot için)."""
    def _write(w):
        pd.DataFrame({"A": [1]}).to_excel(w, sheet_name="Pivot", index=False)
        pd.DataFrame({"A": [1]}).to_excel(w, sheet_name="Detail", index=False)

    out = build_styled_workbook(_write, skip_autofilter_sheets=("Pivot",))
    wb = _load(out)
    assert wb["Pivot"].auto_filter.ref is None
    assert wb["Detail"].auto_filter.ref is not None


# ─────────────────────────────────────────────────────────────────────────────
# TOTAL row highlighting
# ─────────────────────────────────────────────────────────────────────────────
def test_total_row_highlighted_when_first_col_is_total():
    """Birinci kolonu 'TOTAL' olan satır mavi fill + bold ile vurgulanmalı."""
    df = pd.DataFrame({
        "Category": ["A", "B", "TOTAL"],
        "Count":    [3, 5, 8],
    })
    out = build_styled_workbook(
        lambda w: df.to_excel(w, sheet_name="Summary", index=False),
        total_row_sheets=("Summary",),
    )
    ws = _load(out)["Summary"]

    # TOTAL satırı 4. satır (1=header, 2=A, 3=B, 4=TOTAL)
    total_cell = ws.cell(row=4, column=1)
    count_cell = ws.cell(row=4, column=2)
    assert total_cell.value == "TOTAL"
    assert total_cell.font.bold is True
    assert total_cell.fill.fgColor.rgb.upper().endswith(TOTAL_FILL_HEX.upper())
    # Aynı satırın diğer kolonu da vurgulanmalı
    assert count_cell.font.bold is True


def test_total_row_skipped_when_sheet_not_listed():
    """`total_row_sheets` listesinde olmayan sayfada TOTAL satırı normal kalır."""
    df = pd.DataFrame({"X": ["foo", "TOTAL"], "Y": [1, 2]})
    out = build_styled_workbook(
        lambda w: df.to_excel(w, sheet_name="Detail", index=False)
    )
    ws = _load(out)["Detail"]
    total_cell = ws.cell(row=3, column=1)
    assert total_cell.value == "TOTAL"
    # Bold olmamalı (helper bu sayfayı total_row_sheets'e dahil etmedi)
    assert total_cell.font.bold is None or total_cell.font.bold is False


# ─────────────────────────────────────────────────────────────────────────────
# Last-column highlight (Pivot TOTAL kolonu)
# ─────────────────────────────────────────────────────────────────────────────
def test_last_column_highlighted_for_pivot_sheets():
    """Pivot için son kolon (TOTAL) bold + mavi fill almalı."""
    df = pd.DataFrame({
        "Mahalle": ["A", "B"],
        "X":       [1, 2],
        "TOTAL":   [3, 4],
    })
    out = build_styled_workbook(
        lambda w: df.to_excel(w, sheet_name="Pivot", index=False),
        highlight_last_column_sheets=("Pivot",),
    )
    ws = _load(out)["Pivot"]
    # Son kolon 3, 2. ve 3. satır kontrol
    for row_idx in (2, 3):
        c = ws.cell(row=row_idx, column=3)
        assert c.font.bold is True
        assert c.fill.fgColor.rgb.upper().endswith(TOTAL_FILL_HEX.upper())


# ─────────────────────────────────────────────────────────────────────────────
# Column width auto-fit
# ─────────────────────────────────────────────────────────────────────────────
def test_long_value_widens_column_up_to_cap():
    """Uzun değer kolon genişliğini artırmalı, ama DEFAULT_MAX_COL_WIDTH'i aşmamalı."""
    long_str = "x" * 100  # 100 karakter — kapağın üstünde
    out = build_styled_workbook(
        lambda w: pd.DataFrame({"A": [long_str]}).to_excel(
            w, sheet_name="X", index=False
        )
    )
    ws = _load(out)["X"]
    width = ws.column_dimensions["A"].width
    assert width <= DEFAULT_MAX_COL_WIDTH
    # Default min'in (10) üstüne çıkmalı
    assert width > 10


def test_short_value_uses_default_min_width():
    """Çok kısa değer için kolon genişliği default min'de kalmalı."""
    out = build_styled_workbook(
        lambda w: pd.DataFrame({"A": ["x"]}).to_excel(w, sheet_name="X", index=False)
    )
    ws = _load(out)["X"]
    # Header "A" + min width fallback
    assert ws.column_dimensions["A"].width >= 10


# ─────────────────────────────────────────────────────────────────────────────
# Multi-sheet integration (regresyon: 3 sayfa Excel'i tek geçişte stillenmeli)
# ─────────────────────────────────────────────────────────────────────────────
def test_multi_sheet_all_get_header_style():
    """Birden fazla sayfada da header satırı doğru stillenmeli (Optimization 9-sayfa)."""
    def _write(w):
        for name in ("S1", "S2", "S3"):
            pd.DataFrame({"K": [1]}).to_excel(w, sheet_name=name, index=False)

    out = build_styled_workbook(_write)
    wb = _load(out)
    for name in ("S1", "S2", "S3"):
        ws = wb[name]
        assert ws["A1"].font.bold is True
        assert ws.freeze_panes == "A2"


def test_no_autofilter_when_only_header_row():
    """Veri satırı yoksa autofilter set edilmemeli."""
    out = build_styled_workbook(
        lambda w: pd.DataFrame(columns=["A", "B"]).to_excel(
            w, sheet_name="Empty", index=False
        )
    )
    ws = _load(out)["Empty"]
    assert ws.auto_filter.ref is None


# ─────────────────────────────────────────────────────────────────────────────
# style_workbook standalone (write-then-style flow)
# ─────────────────────────────────────────────────────────────────────────────
def test_style_workbook_can_be_used_with_explicit_writer():
    """style_workbook — `pd.ExcelWriter` kapanmadan önce stil uygulayabilmeli."""
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        pd.DataFrame({"A": [1, 2]}).to_excel(writer, sheet_name="X", index=False)
        style_workbook(writer)
    wb = _load(buf.getvalue())
    assert wb["X"]["A1"].font.bold is True
