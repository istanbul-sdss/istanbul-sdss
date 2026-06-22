"""
İndirilebilir şablon üreticileri için testler (Madde 1.2 + 1.3).

Şablonlar UI'da kullanıcıya sunulan boş formatlardır; doğru kolonlara,
geçerli örnek satıra ve doğru encoding'e sahip olmalı — aksi halde
kullanıcı yanlış formatta veri hazırlayıp tool'a yükler ve hata alır.
"""
from __future__ import annotations

from io import BytesIO

import pandas as pd

from src.optimizer.data_loader import (
    POP_METHOD_FOOTPRINT,
    POP_METHOD_UNIFORM,
    load_from_excel,
)
from src.optimizer.templates import (
    build_data_template_xlsx,
    build_tuik_template_xlsx,
)


# ── Data template (Buildings + Assembly) ──────────────────────────────────
def test_data_template_xlsx_has_required_sheets():
    raw = build_data_template_xlsx()
    assert raw and len(raw) > 100, "Şablon boş veya çok küçük"

    xl = pd.ExcelFile(BytesIO(raw))
    assert "Buildings" in xl.sheet_names
    assert "Assembly" in xl.sheet_names
    assert "README" in xl.sheet_names


def test_data_template_xlsx_buildings_has_required_columns():
    raw = build_data_template_xlsx()
    df = pd.read_excel(BytesIO(raw), sheet_name="Buildings")

    must_have = {"Latitude", "Longitude", "Neighbourhood", "Area (m²)", "Floors"}
    assert must_have.issubset(set(df.columns)), (
        f"Buildings sheet eksik kolon: {must_have - set(df.columns)}"
    )
    # Tek örnek satır olmalı
    assert len(df) == 1


def test_data_template_xlsx_assembly_has_required_columns():
    raw = build_data_template_xlsx()
    df = pd.read_excel(BytesIO(raw), sheet_name="Assembly")

    must_have = {"Latitude", "Longitude", "Name", "Area (m²)"}
    assert must_have.issubset(set(df.columns))
    assert len(df) == 1


def test_data_template_xlsx_example_coords_are_istanbul():
    """Örnek koordinatlar İstanbul bounding box içinde olmalı — kullanıcı
    'bu format çalışıyor mu?' diye direkt yükleyebilsin."""
    raw = build_data_template_xlsx()
    bdf = pd.read_excel(BytesIO(raw), sheet_name="Buildings")
    adf = pd.read_excel(BytesIO(raw), sheet_name="Assembly")
    for df_name, df in [("Buildings", bdf), ("Assembly", adf)]:
        lat = float(df["Latitude"].iloc[0])
        lon = float(df["Longitude"].iloc[0])
        assert 40.55 <= lat <= 41.65, f"{df_name} örnek lat={lat} İstanbul dışı"
        assert 27.95 <= lon <= 30.10, f"{df_name} örnek lon={lon} İstanbul dışı"


def test_data_template_xlsx_is_loadable_by_load_from_excel(tmp_path):
    """
    En kritik test: şablonun kendisi `load_from_excel` ile sorunsuz okunabilmeli.
    Bu garanti şablonun ürettiği format ile loader'ın beklediği format arasındaki
    drift'i (örn. kolon adı değişimi) anında yakalar.
    """
    raw = build_data_template_xlsx()
    p = tmp_path / "template.xlsx"
    p.write_bytes(raw)

    b, t = load_from_excel(
        p, "Buildings", "Assembly",
        population_method=POP_METHOD_FOOTPRINT,
    )
    assert len(b) == 1
    assert len(t) == 1
    assert b["weight"].iloc[0] > 0    # footprint-based hesap çalıştı
    assert t["kapasite"].iloc[0] > 0  # kapasite üretildi


# ── TÜİK template (Excel) ─────────────────────────────────────────────────
def test_tuik_template_xlsx_has_required_sheets():
    raw = build_tuik_template_xlsx()
    xl = pd.ExcelFile(BytesIO(raw))
    assert "TUIK_Population" in xl.sheet_names
    assert "README" in xl.sheet_names


def test_tuik_template_xlsx_columns_and_values():
    raw = build_tuik_template_xlsx()
    df = pd.read_excel(BytesIO(raw), sheet_name="TUIK_Population")
    assert "neighbourhood_name" in df.columns
    assert "population" in df.columns
    # En az 3 örnek satır var
    assert len(df) >= 3
    # Tüm nüfus değerleri pozitif tam sayı
    assert (df["population"] > 0).all()


def test_tuik_template_xlsx_has_turkish_chars():
    """Şablon Türkçe karakter içermeli — utf-8 ile doğru yuvarlanır."""
    raw = build_tuik_template_xlsx()
    df = pd.read_excel(BytesIO(raw), sheet_name="TUIK_Population")
    text = " ".join(df["neighbourhood_name"].astype(str).tolist())
    assert "ğ" in text or "ç" in text or "ö" in text, (
        f"Türkçe karakter beklenirdi: {text}"
    )


def test_tuik_template_xlsx_is_styled():
    """build_styled_workbook'tan geçtiği için header stili olmalı."""
    from openpyxl import load_workbook
    raw = build_tuik_template_xlsx()
    wb = load_workbook(BytesIO(raw))
    ws = wb["TUIK_Population"]
    # Header satırı stilize edildi mi (fill veya bold font)?
    header_cell = ws.cell(row=1, column=1)
    is_styled = bool(
        (header_cell.fill and header_cell.fill.fgColor
         and header_cell.fill.fgColor.value not in (None, "00000000"))
        or (header_cell.font and header_cell.font.bold)
    )
    assert is_styled, "TÜİK template header'ında stil bulunamadı"


def test_data_template_xlsx_is_styled():
    """Buildings sheet'inin header'ı styled."""
    from openpyxl import load_workbook
    raw = build_data_template_xlsx()
    wb = load_workbook(BytesIO(raw))
    ws = wb["Buildings"]
    header_cell = ws.cell(row=1, column=1)
    is_styled = bool(
        (header_cell.fill and header_cell.fill.fgColor
         and header_cell.fill.fgColor.value not in (None, "00000000"))
        or (header_cell.font and header_cell.font.bold)
    )
    assert is_styled, "Data template Buildings header'ında stil bulunamadı"


def test_data_template_loadable_with_uniform_when_combined_with_tuik(tmp_path):
    """
    End-to-end senaryo: kullanıcı data template'i indirip 1 satır
    Caferağa binası ekler (footprint olmadan), TÜİK template'i ile
    yükler — uniform mode başarılı çalışmalı.
    """
    # Data template + TÜİK template kombinasyonu
    raw_data = build_data_template_xlsx()
    raw_tuik = build_tuik_template_xlsx()

    p_data = tmp_path / "data.xlsx"
    p_data.write_bytes(raw_data)

    tuik_df = pd.read_excel(BytesIO(raw_tuik))

    b, _ = load_from_excel(
        p_data, "Buildings", "Assembly",
        population_method=POP_METHOD_UNIFORM,
        mahalle_pop=tuik_df,
        pop_name_col="neighbourhood_name",
        pop_value_col="population",
    )
    # Şablonda Caferağa örneği var; TÜİK template'inde Caferağa nüfusu
    # mevcut → 1 bina için weight = TÜİK_nüfus / 1.
    assert len(b) == 1
    # Caferağa için TÜİK template değeri (templates.py'da tanımlı).
    expected = float(tuik_df.loc[
        tuik_df["neighbourhood_name"].str.contains("Caferağa", case=False, na=False),
        "population",
    ].iloc[0])
    assert b["weight"].iloc[0] == expected
    assert b["nufus_kaynak"].iloc[0] == "uniform_per_building"
