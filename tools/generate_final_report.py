"""
generate_final_report.py — Build the INDE final report DRAFT (.docx).

Işık University · Faculty of Natural Sciences and Engineering ·
Department of Industrial Engineering. Mirrors the official template structure
(chapter-based, chapter-prefixed figure/table/equation numbering, UN SDGs,
cost-benefit). Content drawn from the actual codebase; Chapter 5 numeric cells
are placeholders ("—") to be filled by running the Kadıköy optimisation.

Run:  python tools/generate_final_report.py
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "figures"
OUT = Path(r"C:\Users\yusuf\Desktop\INDE4912 Graduation Design Project"
           r"\INDE4902_Final_Report_DRAFT.docx")

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x54, 0x96)
FONT = "Times New Roman"


# ── low-level helpers ────────────────────────────────────────────────────────
def _set_cell_bg(cell, hex_color: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.makeelement(qn("w:shd"), {
        qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): hex_color
    })
    tcpr.append(shd)


def body(doc, text, *, italic=False, bold=False, justify=True, indent=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent is not None:
        p.paragraph_format.left_indent = Pt(indent)
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(12)
    r.italic = italic
    r.bold = bold
    return p


def chapter(doc, text):
    """Chapter heading on a fresh page (section break, next page)."""
    doc.add_section(WD_SECTION.NEW_PAGE)
    h = doc.add_heading(level=1)
    h.paragraph_format.space_before = Pt(6)
    h.paragraph_format.space_after = Pt(12)
    r = h.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(16)
    r.bold = True
    r.font.color.rgb = NAVY
    return h


def section(doc, text):
    h = doc.add_heading(level=2)
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(6)
    r = h.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(13)
    r.bold = True
    r.font.color.rgb = BLUE
    return h


def subsec(doc, text):
    h = doc.add_heading(level=3)
    r = h.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = BLUE
    return h


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(12)


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(12)


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(10)
    r.italic = True


def figure(doc, filename, caption_text, width_in=6.0):
    from docx.shared import Inches
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    run = p.add_run()
    run.add_picture(str(FIG / filename), width=Inches(width_in))
    caption(doc, caption_text)


def equation(doc, text, number):
    """Centered equation text with right-aligned (chapter.n) number via tab."""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    from docx.enum.text import WD_TAB_ALIGNMENT
    # tab stop at right margin (~6.5in content => 9360 twips)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(468), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run("\t" + text + "\t(" + number + ")")
    r.font.name = FONT
    r.font.size = Pt(12)
    r.italic = True


def table(doc, header, rows, widths=None):
    ncols = len(header)
    t = doc.add_table(rows=1, cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(header):
        _set_cell_bg(hdr[i], "1F3864")
        para = hdr[i].paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(htext)
        run.font.name = FONT
        run.font.size = Pt(11)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        if ri % 2 == 1:
            for c in cells:
                _set_cell_bg(c, "EEF2F7")
        for ci, val in enumerate(row):
            para = cells[ci].paragraphs[0]
            if ci > 0:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(str(val))
            run.font.name = FONT
            run.font.size = Pt(11)
    return t


def page_break(doc):
    doc.add_page_break()


def add_toc(doc):
    """Insert a Word TOC field (updates on open with F9)."""
    p = doc.add_paragraph()
    run = p.add_run()
    fldChar = run._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "begin"})
    instr = run._r.makeelement(qn("w:instrText"), {qn("xml:space"): "preserve"})
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar2 = run._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "separate"})
    t = run._r.makeelement(qn("w:t"), {})
    t.text = "Right-click → Update Field to generate the table of contents."
    fldChar3 = run._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "end"})
    run._r.append(fldChar)
    run._r.append(instr)
    run._r.append(fldChar2)
    run._r.append(t)
    run._r.append(fldChar3)


def centered(doc, text, *, size=12, bold=True, after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    return p


# ── document ─────────────────────────────────────────────────────────────────
def build():
    doc = Document()

    # Base style
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)

    sec = doc.sections[0]
    from docx.shared import Cm
    for s in (sec,):
        s.top_margin = Cm(2.5)
        s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(2.5)
        s.right_margin = Cm(2.5)

    # ── TITLE PAGE ──
    for _ in range(2):
        doc.add_paragraph()
    centered(doc, "IŞIK UNIVERSITY", size=14)
    centered(doc, "Faculty of Natural Sciences and Engineering", size=14)
    centered(doc, "Department of Industrial Engineering", size=14, after=36)
    centered(doc, "INDE 4902 — Graduation Design Project", size=13)
    centered(doc, "FINAL REPORT", size=15, after=36)
    centered(doc,
             "Design of a Spatial Decision Support System for "
             "Earthquake Assembly Area Planning", size=16, after=48)
    centered(doc, "by", bold=False, size=12, after=4)
    centered(doc, "Zeliha Karabay — 21INDE1014", bold=False, size=12, after=2)
    centered(doc, "Yusuf Kılıç — 23INDE1049", bold=False, size=12, after=2)
    centered(doc, "Sevcan Görgülü — 21INDE1042", bold=False, size=12, after=36)
    centered(doc, "Supervised by: Seda Baş Güre & İsmail Kayahan",
             bold=False, size=12, after=36)
    centered(doc, "İstanbul — May 2026", size=12)

    # ── DECLARATION ──
    doc.add_section(WD_SECTION.NEW_PAGE)
    section(doc, "Declaration of Academic Integrity")
    body(doc,
         "We hereby declare that this report and the work presented in it are "
         "our own and have been generated by us as the result of our own "
         "original research. Where we have consulted or quoted the work of "
         "others, this is always clearly attributed. With the exception of "
         "such quotations, this report is entirely our own work. We have "
         "acknowledged all main sources of help, and we have made clear "
         "exactly what was contributed by each member of the team.")
    body(doc, "Zeliha Karabay   ______________________")
    body(doc, "Yusuf Kılıç         ______________________")
    body(doc, "Sevcan Görgülü   ______________________")
    body(doc, "Date: ______________________")

    # ── ABSTRACT ──
    doc.add_section(WD_SECTION.NEW_PAGE)
    section(doc, "Abstract")
    body(doc,
         "Istanbul lies adjacent to the North Anatolian Fault and faces a high "
         "probability of a major earthquake, which makes the planning of "
         "post-disaster emergency assembly areas a critical component of urban "
         "resilience. In current practice, the assignment of the population to "
         "assembly areas is frequently driven by administrative convenience "
         "rather than by quantitative accessibility analysis, leaving some "
         "neighbourhoods underserved. This project develops a reproducible "
         "Spatial Decision Support System (SDSS) that selects and assigns "
         "earthquake assembly areas by combining open spatial data with "
         "operations-research methods. Building footprints, candidate assembly "
         "areas, and administrative boundaries are extracted from OpenStreetMap "
         "through the Overpass API; raw and ambiguous tags (notably the generic "
         "“building=yes”) are resolved by a rule-based classification "
         "engine; building-level population is estimated from footprint area, "
         "number of floors, and TÜİK-based occupancy assumptions; and pedestrian "
         "travel times are computed on the real walking network using OSMnx. The "
         "assignment problem is formulated as a capacity-aware p-median model "
         "(Hakimi, 1964) that minimises population-weighted walking time, solved "
         "exactly with integer linear programming (PuLP/CBC) for small instances "
         "and with a K-medoids heuristic for large ones. The methodology is "
         "demonstrated on the Kadıköy district (~6,665 buildings and 154 "
         "candidate assembly areas, ~244,000 estimated residents), while the "
         "system itself is engineered as a general infrastructure applicable to "
         "all 39 districts of Istanbul. The contribution is an end-to-end, "
         "transparent, and extensible decision-support tool for evidence-based "
         "disaster-preparedness planning.")

    # ── ACKNOWLEDGMENTS ──
    doc.add_section(WD_SECTION.NEW_PAGE)
    section(doc, "Acknowledgments")
    body(doc,
         "We would like to express our sincere gratitude to our supervisors, "
         "Seda Baş Güre and İsmail Kayahan, for their continuous guidance and "
         "valuable feedback throughout this project. We also thank the "
         "Department of Industrial Engineering at Işık University for providing "
         "the academic environment that made this work possible, and the "
         "OpenStreetMap community for the open data on which the system relies. "
         "Finally, we thank our families for their unwavering support.")

    # ── CONTENTS ──
    doc.add_section(WD_SECTION.NEW_PAGE)
    section(doc, "Contents")
    add_toc(doc)

    # ════════════════════ CHAPTER 1 ════════════════════
    chapter(doc, "Chapter 1  Introduction")
    section(doc, "1.1  Project Background")
    body(doc,
         "Modern urban planning and disaster preparedness rely heavily on the "
         "availability of accurate, structured, and functionally classified "
         "spatial data. For a megacity such as Istanbul, which faces significant "
         "seismic risk owing to its proximity to the North Anatolian Fault, the "
         "ability to plan for emergencies depends on a detailed understanding of "
         "the urban landscape — buildings, open spaces, and critical "
         "infrastructure. Yet a significant gap exists between the vast amount "
         "of available raw geographic data and the structured information "
         "required for sophisticated engineering analysis and decision-making.")
    body(doc,
         "The primary source of open geographic data is OpenStreetMap (OSM), a "
         "collaborative project that creates a free, editable map of the world. "
         "Although OSM provides an extensive and continuously updated dataset, "
         "its crowd-sourced nature yields data that are often inconsistent, "
         "incomplete, and ambiguously classified. A prominent example is the "
         "widespread “building=yes” tag, which confirms that a "
         "structure exists but says nothing about its function. This project "
         "turns such raw data into a decision-ready resource and uses it to "
         "optimise earthquake assembly-area planning.")
    section(doc, "1.2  Motivation and Problem Statement")
    body(doc,
         "In the immediate aftermath of an earthquake, residents must reach safe "
         "open spaces — emergency assembly areas — on foot, often under degraded "
         "conditions where vehicular transport is unavailable. The Disaster and "
         "Emergency Management Presidency (AFAD) defines standards for these "
         "areas, including a minimum usable area per person. In practice, "
         "however, deciding which open space serves which residents is frequently "
         "based on administrative boundaries and local knowledge rather than a "
         "systematic accessibility analysis. The result is spatially uneven "
         "provision: dense neighbourhoods may be served by distant or undersized "
         "areas, and the true walking accessibility of the population is rarely "
         "quantified. The problem addressed by this project is the selection and "
         "capacity-aware assignment of assembly areas so that the population is "
         "served within acceptable walking times.")
    section(doc, "1.3  Project Objectives")
    body(doc, "The specific objectives of the project are:")
    numbered(doc, "To extract and clean building, assembly-area, and "
                  "administrative-boundary data for Istanbul districts from "
                  "OpenStreetMap, resolving ambiguous tags through a rule-based "
                  "classification engine.")
    numbered(doc, "To estimate building-level population as the demand weight "
                  "for the optimisation.")
    numbered(doc, "To compute realistic pedestrian travel times between "
                  "buildings and assembly areas using the street network.")
    numbered(doc, "To formulate and solve a capacity-aware p-median model, "
                  "providing both an exact solver and a scalable heuristic.")
    numbered(doc, "To deliver the methodology as an interactive decision-support "
                  "tool with maps, key indicators, and exportable reports.")
    section(doc, "1.4  Relevance to United Nations SDGs")
    body(doc,
         "The project contributes directly to several United Nations Sustainable "
         "Development Goals. It supports SDG 11 (Sustainable Cities and "
         "Communities), and in particular Target 11.5 on reducing deaths and the "
         "number of people affected by disasters, by improving the spatial "
         "planning of emergency assembly areas. It supports SDG 9 (Industry, "
         "Innovation and Infrastructure) by building a reusable, open-data "
         "analytical infrastructure. It also contributes to SDG 13 (Climate "
         "Action), since the same accessibility and capacity analysis applies to "
         "preparedness for climate-related hazards that require population "
         "evacuation to safe areas.")
    section(doc, "1.5  Contributions")
    body(doc, "The main contributions of this project are:")
    bullet(doc, "An automated, open-source pipeline that converts raw OSM data "
                "into functionally classified, analysis-ready spatial datasets "
                "for any Istanbul district.")
    bullet(doc, "A capacity-aware p-median optimisation model with both an exact "
                "(ILP) and a heuristic (K-medoids) solver, exposing efficiency "
                "and equity objectives.")
    bullet(doc, "An interactive decision-support application that makes the "
                "methodology, its assumptions, and its solution quality "
                "transparent to non-specialist planners.")
    bullet(doc, "A demonstrated case study on the Kadıköy district, with a "
                "design that generalises to all 39 districts of Istanbul.")

    # ════════════════════ CHAPTER 2 ════════════════════
    chapter(doc, "Chapter 2  Problem Description")
    section(doc, "2.1  Detailed Problem Definition")
    body(doc,
         "Istanbul is home to more than fifteen million people and sits on one "
         "of the most dangerous seismic zones in the world. Scientists largely "
         "agree that a major earthquake affecting the Marmara region is likely "
         "within the coming decades. When such an earthquake strikes, the first "
         "minutes and hours are decisive: people leave damaged buildings and "
         "need to gather in safe, open places where they can wait in safety, be "
         "counted, receive first aid, and be reached by rescue teams. These safe "
         "open places are called emergency assembly areas. How well a city has "
         "planned these areas — how many there are, where they are, and how "
         "easily people can walk to them — has a direct effect on how many lives "
         "can be protected in a disaster.")
    body(doc,
         "Today, the way assembly areas are chosen and assigned to residents is "
         "mostly based on administrative habit rather than on measurement. A "
         "park or a school yard is often declared an assembly area simply "
         "because it is well known or because it falls inside a particular "
         "neighbourhood, without checking whether the people who are supposed to "
         "use it can actually reach it quickly on foot, or whether it is large "
         "enough to hold them all. As a result, some crowded neighbourhoods are "
         "served by assembly areas that are too far away or too small, while the "
         "real walking accessibility of the population is almost never "
         "calculated. In an emergency, these hidden weaknesses can turn into "
         "real danger.")
    body(doc,
         "There is also a data problem behind the planning problem. To analyse "
         "accessibility, a planner first needs to know where the buildings are, "
         "how many people live in them, and where the usable open spaces are. "
         "This information is, in principle, available for free on "
         "OpenStreetMap, a kind of “Wikipedia of maps” that volunteers "
         "around the world keep up to date. Unfortunately, because anyone can "
         "contribute, the data are messy. Many buildings are recorded only as "
         "“a building”, with no indication of whether they are homes, "
         "schools, hospitals, or shops. Different contributors describe the same "
         "kind of place in different ways, and some information is missing "
         "altogether. Before any meaningful planning can take place, this raw "
         "and inconsistent data has to be cleaned, completed, and organised — a "
         "task that is far too large to do by hand for a whole city.")
    body(doc,
         "This project tackles both problems together. First, it builds a "
         "software system that automatically collects the map data for a chosen "
         "district, cleans it, and fills in the missing information about what "
         "each building is used for, so that the city can be described reliably "
         "in terms of where people are and where they can shelter. Second, it "
         "uses this clean information to answer the planning question directly: "
         "given a limited number of assembly areas that can realistically be "
         "opened and equipped, which ones should be chosen, and which residents "
         "should be sent to each one, so that everybody can reach safety as "
         "quickly as possible without any area being overcrowded? The system "
         "treats this as a precise optimisation question and finds the best "
         "answer it can, instead of relying on guesswork.")
    body(doc,
         "The value of solving this problem is easy to understand even for "
         "someone with no technical background. Imagine two plans for the same "
         "district. In the first, drawn up by habit, a family living in a dense "
         "back street has to walk twenty minutes through narrow roads to reach a "
         "crowded park. In the second, produced by the system in this project, "
         "the same family is assigned to a nearer open space that has enough room "
         "for everyone in the surrounding blocks, and their walk is cut to a few "
         "minutes. Multiplied across hundreds of thousands of residents, this "
         "difference is measured in safety and, ultimately, in lives. The "
         "purpose of the project is to give planners a trustworthy, transparent "
         "tool that turns the city's own open data into exactly this kind of "
         "improved plan, and to show that the same tool can be applied to every "
         "district of Istanbul, not just the one studied in detail here.")
    body(doc,
         "It is important to be clear about what the project does and does not "
         "claim. It does not replace the official authorities or produce a "
         "legally binding assembly-area plan; that would require institutional "
         "data and formal approval. Instead, it provides a rigorous, repeatable "
         "method and a working tool that authorities, researchers, and students "
         "can use to evaluate and improve such plans. Its strength is that every "
         "step — where the data came from, how the population was estimated, how "
         "walking times were measured, and how the final assignment was chosen — "
         "is explicit and can be checked, so that the resulting recommendations "
         "can be trusted and defended.")
    section(doc, "2.2  Scope and Assumptions")
    body(doc,
         "The optimisation is demonstrated in depth on the Kadıköy district, "
         "which serves as the representative case study throughout the report. "
         "The supporting system, however, is designed as a general "
         "infrastructure: the data-extraction, classification, "
         "population-estimation, routing, and optimisation modules all operate "
         "on any of Istanbul's 39 districts without modification. The main "
         "assumptions are:")
    bullet(doc, "Population is estimated from building footprint, number of "
                "floors, and a TÜİK-based occupancy factor; it is not measured "
                "at the individual level.")
    bullet(doc, "Walking time is computed on the pedestrian street network at a "
                "constant speed of 4.8 km/h on level ground; slope, stairs, and "
                "post-disaster obstructions are not modelled.")
    bullet(doc, "Assembly-area capacity is derived from usable area divided by a "
                "density of 1.5 m²/person (the AFAD assembly standard, "
                "adjustable by the user).")
    bullet(doc, "Candidate assembly areas are taken from OSM tags (parks, school "
                "yards, open spaces) rather than from an official register.")
    section(doc, "2.3  Limitations and Restrictions")
    body(doc,
         "The project is bounded by the quality of open data and by the absence "
         "of official institutional datasets. OSM attributes are "
         "volunteer-contributed and may be missing or inconsistent; the "
         "population figures are estimates; and the candidate assembly areas are "
         "inferred rather than certified. The model is deterministic and does "
         "not represent demand uncertainty or scenario-based network damage. "
         "Computational resources also bound the exact solver, which is why a "
         "heuristic is provided for the largest instances. These restrictions "
         "define the interpretation of the results rather than invalidate the "
         "methodology.")

    # ════════════════════ CHAPTER 3 ════════════════════
    chapter(doc, "Chapter 3  Literature Review")
    section(doc, "3.1  Review of Relevant Studies")
    body(doc,
         "This project draws on three established research areas: facility "
         "location in operations research, location modelling for disaster "
         "management, and the use of open spatial data with network analysis.")
    body(doc,
         "Facility location and the p-median problem. The p-median problem, "
         "introduced by Hakimi (1964, 1965), seeks the locations of p facilities "
         "that minimise the total demand-weighted distance between demand points "
         "and their nearest facility. ReVelle and Swain (1970) gave an integer "
         "linear-programming formulation and showed that strong solutions could "
         "be obtained via linear relaxation, establishing the computational "
         "tradition this project follows. Comprehensive treatments of discrete "
         "location models are provided by Daskin (2013). Related models encode "
         "different equity–efficiency trade-offs: the set-covering and "
         "maximal-covering problems (Toregas et al., 1971; Church and ReVelle, "
         "1974) optimise demand covered within a fixed threshold, while the "
         "p-center problem minimises the worst-case distance. This project "
         "adopts the p-median objective for efficiency but additionally reports "
         "worst-case and 95th-percentile travel times and offers "
         "fairness-oriented objectives.")
    body(doc,
         "Location models in disaster management. Facility-location models are "
         "widely applied to emergency facilities such as ambulance and fire "
         "stations, shelters, and assembly areas, where coverage and "
         "accessibility within a critical time window are paramount. For "
         "post-earthquake assembly and shelter planning, Anhorn and Khazai "
         "(2015) propose an open-space suitability analysis for emergency "
         "shelter allocation. A consistent message of this literature is the "
         "importance of realistic, network-based travel distances and of "
         "incorporating capacity — both central to the present work.")
    body(doc,
         "Open spatial data and network analysis. Volunteered geographic "
         "information, principally OpenStreetMap, has made fine-grained urban "
         "analysis feasible without proprietary data (Goodchild, 2007; Haklay "
         "and Weber, 2008). Boeing (2017) introduced OSMnx, a Python library for "
         "downloading and analysing street networks from OSM, enabling "
         "reproducible network-based accessibility computation. This project "
         "uses OSM as its sole spatial source and OSMnx for pedestrian routing.")
    section(doc, "3.2  Summary and Taxonomy")
    body(doc,
         "Table 3.1 positions the present work against representative studies "
         "according to the modelling objective, the treatment of capacity, the "
         "distance model, and the data source. While each ingredient is "
         "individually well established, integrated and reusable tools that "
         "combine them for district-scale earthquake assembly planning with an "
         "accessible interface remain scarce — the gap this project addresses.")
    table(doc,
          ["Study", "Objective", "Capacity", "Distance", "Data"],
          [
              ["Hakimi (1964)", "p-median", "No", "Network", "Generic"],
              ["ReVelle & Swain (1970)", "p-median (ILP)", "No", "Network", "Generic"],
              ["Church & ReVelle (1974)", "Max covering", "No", "Threshold", "Generic"],
              ["Anhorn & Khazai (2015)", "Suitability/coverage", "Partial", "Network", "GIS"],
              ["This project", "p-median + fairness", "Yes", "Network (walk)", "OSM / open"],
          ])
    caption(doc, "Table 3.1  Taxonomy of reviewed studies versus the present work.")

    # ════════════════════ CHAPTER 4 ════════════════════
    chapter(doc, "Chapter 4  System Design and Methodology")
    section(doc, "4.1  Proposed Solution Approach")
    body(doc,
         "The selected approach is mathematical optimisation, supported by a "
         "spatial-data engineering pipeline. Optimisation was preferred over "
         "pure simulation or heuristic-only design because the assignment of "
         "buildings to assembly areas has a clear objective (minimise "
         "population-weighted walking time), well-defined constraints (each "
         "building served once, a fixed number of areas opened, capacity "
         "respected), and a need for solutions whose quality can be stated "
         "precisely. The system is organised as a modular pipeline with four "
         "stages — data extraction, preprocessing and population estimation, "
         "origin–destination (OD) travel-time computation, and p-median "
         "optimisation — followed by visualisation and reporting. Figure 4.1 "
         "shows the system architecture and Figure 4.2 the optimisation "
         "workflow.")
    figure(doc, "figure_4_1_sdss_architecture.png",
           "Figure 4.1  High-level architecture of the spatial decision-support system.")
    figure(doc, "figure_4_2_optimization_workflow.png",
           "Figure 4.2  Optimisation workflow from data loading to results.")
    section(doc, "4.2  System Design / Mathematical Modeling")
    body(doc,
         "The assignment of buildings to assembly areas is formulated as a "
         "capacity-aware p-median model. The sets, parameters, decision "
         "variables, objective, and constraints are defined below.")
    body(doc, "Sets and indices:", bold=True)
    bullet(doc, "I — set of demand points (buildings), indexed by i.")
    bullet(doc, "J — set of candidate facilities (assembly areas), indexed by j.")
    body(doc, "Parameters:", bold=True)
    bullet(doc, "w_i — estimated population of building i (persons).")
    bullet(doc, "d_ij — walking time from building i to area j (minutes).")
    bullet(doc, "C_j — capacity of area j (persons) = usable area (m²) / density (m²/person).")
    bullet(doc, "p — number of assembly areas to open (count).")
    body(doc, "Decision variables:", bold=True)
    bullet(doc, "y_j ∈ {0,1} — 1 if area j is opened, 0 otherwise.")
    bullet(doc, "x_ij ∈ {0,1} — 1 if building i is assigned to area j, 0 otherwise.")
    body(doc, "Objective function (efficiency, min-sum):", bold=True)
    equation(doc, "minimise  Σ_i Σ_j  w_i · d_ij · x_ij", "4.1")
    body(doc, "Constraints:", bold=True)
    equation(doc, "Σ_j x_ij = 1                 for all i ∈ I", "4.2")
    body(doc, "Each building is assigned to exactly one assembly area.")
    equation(doc, "x_ij ≤ y_j                   for all i ∈ I, j ∈ J", "4.3")
    body(doc, "A building may be assigned only to an opened area.")
    equation(doc, "Σ_j y_j = p", "4.4")
    body(doc, "Exactly p assembly areas are opened.")
    equation(doc, "Σ_i w_i · x_ij ≤ C_j · y_j    for all j ∈ J", "4.5")
    body(doc, "Population assigned to an area does not exceed its capacity "
              "(enforced when the capacity constraint is enabled).")
    equation(doc, "x_ij, y_j ∈ {0,1}", "4.6")
    body(doc,
         "Two alternative objectives are provided to expose the equity "
         "dimension: a min-max objective that minimises the single worst "
         "assignment time, and a population-weighted 95th-percentile objective "
         "that is robust to outliers and is recommended for decision support. "
         "Capacity C_j is derived from each area's usable area divided by a "
         "density parameter (default 1.5 m²/person, the AFAD standard), which is "
         "exposed to the user for sensitivity analysis.")
    section(doc, "4.3  Data Collection and Preparation")
    body(doc,
         "Spatial data are collected from OpenStreetMap via the Overpass API. "
         "The extraction module retrieves building footprints, points of "
         "interest (including candidate assembly areas such as parks and school "
         "yards), and district and neighbourhood boundaries, organised by a rule "
         "registry of nine main categories and more than one hundred "
         "subcategories. Each record is deduplicated, geocoded to its "
         "neighbourhood, and assigned a confidence score. A rule-based "
         "classification engine resolves ambiguous tags — for example, a "
         "“building=yes” feature whose name contains "
         "“Hastanesi” is classified as a healthcare facility — using "
         "existing tags, the feature name, and spatial relationships. Queries "
         "are retried across multiple Overpass mirrors, and partial or total "
         "failures are reported explicitly rather than silently treated as "
         "“no data”. Population is then estimated at the building "
         "level (footprint × floors × occupancy factor, or a uniform "
         "neighbourhood distribution), and geometries are projected to UTM zone "
         "35N for accurate area and distance computation.")
    section(doc, "4.4  Solution Method")
    body(doc,
         "Walking times are computed on the pedestrian network extracted via "
         "OSMnx: the district graph is downloaded once and cached, each building "
         "and area is mapped to its nearest network node, and shortest-path "
         "times are computed from each facility to all demand nodes at 4.8 km/h. "
         "Where the network cannot be retrieved, a haversine fallback with a "
         "detour factor of 1.4 produces a clearly-labelled estimate. Problem "
         "size determines the solver: at or below 5,000 buildings the model is "
         "solved exactly as an integer linear program using PuLP with the CBC "
         "solver, yielding a provably optimal assignment; above the threshold a "
         "capacity-aware K-medoids heuristic with greedy initialisation and "
         "1-swap local search is used, returning a high-quality approximate "
         "solution and reporting whether the local search converged or hit its "
         "iteration limit. The tool can be forced into either mode for "
         "comparison and can use faster ILP back-ends (e.g., HiGHS) when "
         "installed.")

    # ════════════════════ CHAPTER 5 ════════════════════
    chapter(doc, "Chapter 5  Computational Results and Analysis")
    body(doc,
         "[NOTE TO AUTHORS: the numeric cells marked “—” are produced "
         "by running the Kadıköy optimisation described below on the cached "
         "dataset (data/samples/Kadıköy_OSM_sample.xlsx + the cached walking "
         "graph). Replace them with the tool's output, and insert the coverage "
         "curve, assignment map, and K-medoids convergence chart exported from "
         "the optimisation tool.]", italic=True)
    section(doc, "5.1  Experimental Design / Scenarios")
    body(doc,
         "Experiments use the Kadıköy dataset: approximately 6,665 buildings "
         "(demand) and 154 candidate assembly areas (facilities), with a total "
         "estimated population of about 244,000. Because this instance sits just "
         "above the exact/heuristic threshold, it is well suited to comparing "
         "the two solver modes. Five analyses are performed: (i) a base "
         "optimisation under the efficiency objective; (ii) a sensitivity "
         "analysis over the number of opened areas p; (iii) a comparison of "
         "capacity-on versus capacity-off; (iv) a density sensitivity sweep over "
         "the area-per-person parameter; and (v) a comparison of the exact (ILP) "
         "and heuristic (K-medoids) solvers in objective value and runtime.")
    section(doc, "5.2  Analysis of Results")
    body(doc, "Base optimisation. Table 5.1 summarises the base optimisation "
              "under the efficiency objective.")
    table(doc,
          ["Indicator", "Value", "Unit"],
          [
              ["Opened assembly areas (p)", "—", "count"],
              ["Population-weighted avg. time", "—", "minutes"],
              ["Worst-case time", "—", "minutes"],
              ["Population-weighted p95 time", "—", "minutes"],
              ["Population covered ≤ 5 min", "—", "%"],
              ["Population covered ≤ 10 min", "—", "%"],
              ["Solver / status", "—", "ILP / K-medoids"],
              ["Runtime", "—", "seconds"],
          ])
    caption(doc, "Table 5.1  Base optimisation results for Kadıköy (efficiency objective).")
    body(doc, "Sensitivity to the number of opened areas. Table 5.2 shows how "
              "coverage and travel time change as p increases.")
    table(doc,
          ["p", "Avg time (min)", "p95 (min)", "≤5 min (%)", "≤10 min (%)"],
          [["3", "—", "—", "—", "—"],
           ["5", "—", "—", "—", "—"],
           ["10", "—", "—", "—", "—"],
           ["15", "—", "—", "—", "—"],
           ["20", "—", "—", "—", "—"]])
    caption(doc, "Table 5.2  Sensitivity of coverage and travel time to the number of opened areas.")
    body(doc, "Effect of capacity. Table 5.3 compares the solution with the "
              "capacity constraint disabled and enabled.")
    table(doc,
          ["Indicator", "Capacity OFF", "Capacity ON"],
          [["Population-weighted avg. time (min)", "—", "—"],
           ["p95 time (min)", "—", "—"],
           ["Population covered ≤ 10 min (%)", "—", "—"],
           ["Max. area utilisation (%)", "—", "—"]])
    caption(doc, "Table 5.3  Capacity-on versus capacity-off comparison.")
    body(doc, "Density sensitivity. Table 5.4 sweeps the area-per-person density "
              "across the emergency (1.0), AFAD assembly (1.5), and long-term "
              "shelter (2.5) standards.")
    table(doc,
          ["Density (m²/person)", "Feasible?", "Avg time (min)"],
          [["1.0", "—", "—"], ["1.5", "—", "—"], ["2.5", "—", "—"]])
    caption(doc, "Table 5.4  Sensitivity to the area-per-person density parameter.")
    body(doc, "Exact versus heuristic. Table 5.5 compares the ILP and K-medoids "
              "solvers on the same instance.")
    table(doc,
          ["Metric", "Exact (ILP)", "Heuristic (K-medoids)"],
          [["Objective (weighted minutes)", "—", "—"],
           ["Optimality gap (%)", "0 (reference)", "—"],
           ["Runtime (s)", "—", "—"],
           ["Converged / status", "—", "—"]])
    caption(doc, "Table 5.5  Exact versus heuristic solver comparison on the Kadıköy instance.")
    section(doc, "5.3  Limitations")
    body(doc,
         "The results inherit the limitations of the input data and modelling "
         "assumptions: volunteer-contributed OSM attributes, building-level "
         "population estimates, and a constant walking speed on level ground. "
         "They should therefore be read as a rigorous comparative analysis of "
         "planning options rather than as an officially validated assembly-area "
         "plan.")

    # ════════════════════ CHAPTER 6 ════════════════════
    chapter(doc, "Chapter 6  Conclusions and Discussions")
    section(doc, "6.1  Cost-Benefit Analysis")
    body(doc,
         "The economic case for the system rests on the very low cost of its "
         "inputs against the high value of improved preparedness. The system "
         "uses only open data (OpenStreetMap) and open-source software (Python, "
         "Streamlit, OSMnx, PuLP/CBC), so there are no licensing or data "
         "acquisition costs; the principal cost is the development effort, which "
         "is non-recurring, plus negligible compute. Against this, a "
         "conventional consultancy-led accessibility study for a single district "
         "typically requires paid GIS data and specialist labour. Because the "
         "tool generalises to all 39 districts at essentially zero marginal "
         "cost, the cost per district falls rapidly as it is reused. The benefit "
         "side — better-located assembly areas and shorter, capacity-feasible "
         "evacuation walks — is difficult to monetise directly but is realised "
         "as reduced risk to life and faster post-disaster response, which "
         "dominates any reasonable accounting of the modest development cost.")
    section(doc, "6.2  Economical, Social, Ethical and Environmental Impacts")
    body(doc,
         "Economically, the tool lowers the barrier to evidence-based "
         "preparedness planning for municipalities with limited budgets. "
         "Socially, by quantifying walking accessibility it exposes and helps "
         "correct spatial inequities in emergency provision, and its fairness "
         "objectives let planners protect the worst-served residents rather than "
         "only the average. Ethically, the system is built for transparency: "
         "every assumption and the quality of the underlying data are reported, "
         "user- and OSM-derived text is sanitised before display, and the "
         "distinction between provably optimal and heuristic solutions is made "
         "explicit, so recommendations are not overstated. Environmentally, the "
         "approach relies on existing open data and lightweight computation, "
         "avoiding new field surveys, and it promotes the protective use of "
         "existing urban open and green spaces.")
    section(doc, "6.3  Conclusions and Future Work")
    body(doc,
         "This project designed and implemented a reproducible spatial "
         "decision-support system for earthquake assembly-area selection in "
         "Istanbul. It integrates open-data extraction and rule-based "
         "classification, building-level population estimation, network-based "
         "pedestrian travel-time computation, and a capacity-aware p-median "
         "optimisation offering both an exact solver and a scalable heuristic, "
         "delivered as interactive applications with maps, key indicators, and "
         "exportable reports. The methodology was demonstrated on the Kadıköy "
         "district while being engineered as a general infrastructure for all 39 "
         "districts of Istanbul.")
    body(doc,
         "Future work includes validating the output against official AFAD "
         "assembly-area designations; exercising scalability on the largest "
         "districts and ultimately a city-wide run using the heuristic solver "
         "and faster ILP back-ends; incorporating finer demographic data and "
         "scenario-based network damage to represent uncertainty; and unifying "
         "the two applications into a single multi-page, bilingual "
         "(Turkish/English) interface to broaden accessibility for local "
         "planners.")

    # ── BIBLIOGRAPHY ──
    chapter(doc, "Bibliography")
    refs = [
        "Anhorn, J., & Khazai, B. (2015). Open space suitability analysis for "
        "emergency shelter after an earthquake. Natural Hazards and Earth "
        "System Sciences, 15(4), 789–803.",
        "Boeing, G. (2017). OSMnx: New methods for acquiring, constructing, "
        "analyzing, and visualizing complex street networks. Computers, "
        "Environment and Urban Systems, 65, 126–139.",
        "Church, R., & ReVelle, C. (1974). The maximal covering location "
        "problem. Papers in Regional Science, 32(1), 101–118.",
        "Daskin, M. S. (2013). Network and Discrete Location: Models, "
        "Algorithms, and Applications (2nd ed.). Hoboken, NJ: Wiley.",
        "Goodchild, M. F. (2007). Citizens as sensors: the world of "
        "volunteered geography. GeoJournal, 69(4), 211–221.",
        "Haklay, M., & Weber, P. (2008). OpenStreetMap: User-generated street "
        "maps. IEEE Pervasive Computing, 7(4), 12–18.",
        "Hakimi, S. L. (1964). Optimum locations of switching centers and the "
        "absolute centers and medians of a graph. Operations Research, 12(3), "
        "450–459.",
        "Hakimi, S. L. (1965). Optimum distribution of switching centers in a "
        "communication network and some related graph-theoretic problems. "
        "Operations Research, 13(3), 462–475.",
        "ReVelle, C. S., & Swain, R. W. (1970). Central facilities location. "
        "Geographical Analysis, 2(1), 30–42.",
        "Toregas, C., Swain, R., ReVelle, C., & Bergman, L. (1971). The "
        "location of emergency service facilities. Operations Research, 19(6), "
        "1363–1373.",
        "Turkish Statistical Institute (TÜİK). Population and demographic "
        "statistics. Ankara, Türkiye.",
        "Disaster and Emergency Management Presidency (AFAD). Assembly area "
        "standards and guidelines. Ankara, Türkiye.",
        "OpenStreetMap contributors. (2024). OpenStreetMap. "
        "https://www.openstreetmap.org",
    ]
    for r in refs:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.left_indent = Pt(24)
        p.paragraph_format.first_line_indent = Pt(-24)
        p.paragraph_format.space_after = Pt(6)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(r)
        run.font.name = FONT
        run.font.size = Pt(12)

    # ── APPENDIX ──
    chapter(doc, "Appendix")
    body(doc, "Appendix A — Sample data-extraction outputs: a representative "
              "multi-sheet Excel workbook produced by the data-collection tool "
              "for Kadıköy (Summary, Neighbourhood Pivot, per-category detail, "
              "and Data Quality sheets).")
    body(doc, "Appendix B — Optimisation result tables: the full building→area "
              "assignment table and the per-area utilisation summary exported by "
              "the optimisation tool, together with the sensitivity and "
              "comparison results underlying Chapter 5.")
    body(doc, "Appendix C — User-interface screenshots: the data-extraction "
              "page, the interactive map, the analytics dashboard, and the "
              "four-step optimisation workflow.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"WROTE {OUT} ({OUT.stat().st_size:,} bytes)")


if __name__ == "__main__":
    build()
