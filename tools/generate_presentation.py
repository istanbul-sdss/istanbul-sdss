"""
generate_presentation.py — Build the INDE graduation-project presentation (.pptx).

Design system: "Seismic Slate" — deep navy primary, seismic-amber accent,
teal secondary, light surfaces. Cambria headings + Calibri body. 16:9.
Dark title/closing slides, light content slides (sandwich structure),
repeated motif: numbered chip + colored-circle icons, card blocks, stat
callouts. Embeds the two architecture figures; leaves clearly-labelled
placeholders where screenshots/result charts are still to be added.

Run:  python tools/generate_presentation.py
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "figures"
OUT = Path(r"C:\Users\yusuf\Desktop\INDE4912 Graduation Design Project"
           r"\INDE4902_Presentation.pptx")

# ── palette ──────────────────────────────────────────────────────────────────
NAVY      = RGBColor(0x0B, 0x22, 0x36)   # primary dark
NAVY2     = RGBColor(0x12, 0x30, 0x4B)   # surface dark
AMBER     = RGBColor(0xF4, 0x73, 0x2A)   # seismic accent
TEAL      = RGBColor(0x2E, 0x8C, 0x8C)   # secondary
ICE       = RGBColor(0xCA, 0xDC, 0xFC)   # light ice (on dark)
LIGHT     = RGBColor(0xF5, 0xF8, 0xFC)   # page bg (content)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
CARD      = RGBColor(0xFF, 0xFF, 0xFF)
INK       = RGBColor(0x0F, 0x1B, 0x2A)   # body text
MUTED     = RGBColor(0x5A, 0x6B, 0x7B)   # secondary text
HAIR      = RGBColor(0xDD, 0xE5, 0xEE)   # hairline / card border
GREEN     = RGBColor(0x2E, 0x7D, 0x4F)

HEAD = "Cambria"
BODY = "Calibri"

EMU_IN = 914400
SW, SH = 13.333, 7.5


# ── helpers ──────────────────────────────────────────────────────────────────
def _in(v):
    return Emu(int(v * EMU_IN))


def slide(prs, bg=LIGHT):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    return s


def rect(s, x, y, w, h, fill=None, line=None, line_w=1.0, rounded=False, shadow=False):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        _in(x), _in(y), _in(w), _in(h))
    if rounded:
        try:
            shp.adjustments[0] = 0.06
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    if shadow:
        _soft_shadow(shp)
    return shp


def _soft_shadow(shp):
    spPr = shp._element.spPr
    # Reuse an existing (possibly empty) <a:effectLst> instead of adding a
    # second one — two effectLst siblings in one spPr is invalid OOXML and
    # makes PowerPoint flag the file as corrupt.
    el = spPr.find(qn('a:effectLst'))
    if el is None:
        el = spPr.makeelement(qn('a:effectLst'), {})
        spPr.append(el)
    else:
        for child in list(el):
            el.remove(child)
    sh = spPr.makeelement(qn('a:outerShdw'),
                          {'blurRad': '90000', 'dist': '38100',
                           'dir': '5400000', 'rotWithShape': '0'})
    clr = spPr.makeelement(qn('a:srgbClr'), {'val': '0B2236'})
    alpha = spPr.makeelement(qn('a:alpha'), {'val': '20000'})
    clr.append(alpha); sh.append(clr); el.append(sh)


def oval(s, x, y, d, fill, line=None):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, _in(x), _in(y), _in(d), _in(d))
    o.fill.solid(); o.fill.fore_color.rgb = fill
    if line is None:
        o.line.fill.background()
    else:
        o.line.color.rgb = line; o.line.width = Pt(1.25)
    o.shadow.inherit = False
    return o


def text(s, x, y, w, h, runs, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=4, line=1.0, wrap=True):
    """runs: list of (string, size, color, bold, italic, font) tuples or list of
    paragraphs where each paragraph is such a list."""
    tb = s.shapes.add_textbox(_in(x), _in(y), _in(w), _in(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for m in (tf.margin_left, ):
        pass
    tf.margin_left = 0; tf.margin_right = 0
    tf.margin_top = 0; tf.margin_bottom = 0
    # normalize: list of paragraphs
    if runs and isinstance(runs[0], tuple):
        paragraphs = [runs]
    else:
        paragraphs = runs
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        try:
            p.line_spacing = line
        except Exception:
            pass
        for (txt, size, color, bold, italic, font) in para:
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
            r.font.name = font; r.font.color.rgb = color
    return tb


def R(t, size, color, bold=False, italic=False, font=BODY):
    return (t, size, color, bold, italic, font)


def bullet_block(s, x, y, w, items, *, size=15, color=INK, gap=8,
                 marker_color=AMBER, lh=1.04):
    tb = s.shapes.add_textbox(_in(x), _in(y), _in(w), _in(0.5))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.space_before = Pt(0); p.line_spacing = lh
        r0 = p.add_run(); r0.text = "▪  "
        r0.font.size = Pt(size); r0.font.bold = True; r0.font.name = BODY
        r0.font.color.rgb = marker_color
        if isinstance(it, tuple):
            head, rest = it
            r1 = p.add_run(); r1.text = head
            r1.font.size = Pt(size); r1.font.bold = True; r1.font.name = BODY
            r1.font.color.rgb = color
            r2 = p.add_run(); r2.text = rest
            r2.font.size = Pt(size); r2.font.name = BODY; r2.font.color.rgb = color
        else:
            r1 = p.add_run(); r1.text = it
            r1.font.size = Pt(size); r1.font.name = BODY; r1.font.color.rgb = color
    return tb


def kicker(s, x, y, txt, color=AMBER, on_dark=False):
    text(s, x, y, 8, 0.3, [R(txt.upper(), 12.5, color, bold=True, font=BODY)])


def title(s, x, y, w, txt, color=NAVY, size=33):
    text(s, x, y, w, 1.1, [R(txt, size, color, bold=True, font=HEAD)], line=1.0)


def footer(s, idx, total, label, dark=False):
    c = ICE if dark else MUTED
    text(s, 0.55, 7.06, 9, 0.3,
         [R("Işık University · Industrial Engineering · INDE 4902", 9.5, c, font=BODY)])
    text(s, 11.3, 7.06, 1.5, 0.3,
         [R(f"{idx:02d} / {total:02d}", 9.5, c, bold=True, font=BODY)],
         align=PP_ALIGN.RIGHT)


def chip(s, x, y, n, color=AMBER):
    rect(s, x, y, 0.42, 0.42, fill=color, rounded=True)
    text(s, x, y + 0.015, 0.42, 0.42, [R(str(n), 17, WHITE, bold=True, font=HEAD)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def accent_side(s, color=AMBER):
    rect(s, 0, 0, 0.16, SH, fill=color)


TOTAL = 19


def card(s, x, y, w, h, *, fill=CARD, border=HAIR, shadow=True):
    return rect(s, x, y, w, h, fill=fill, line=border, line_w=1.0,
                rounded=True, shadow=shadow)


def placeholder(s, x, y, w, h, label):
    rect(s, x, y, w, h, fill=RGBColor(0xED, 0xF2, 0xF8),
         line=TEAL, line_w=1.25, rounded=True)
    # dashed feel via icon + text
    text(s, x, y + h/2 - 0.55, w, 0.4,
         [R("🖼", 30, TEAL, font=BODY)], align=PP_ALIGN.CENTER)
    text(s, x + 0.3, y + h/2 + 0.0, w - 0.6, 0.9,
         [[R("IMAGE PLACEHOLDER", 12, TEAL, bold=True, font=BODY)],
          [R(label, 11.5, MUTED, italic=True, font=BODY)]],
         align=PP_ALIGN.CENTER, line=1.02, space_after=2)


# ══════════════════════════════════════════════════════════════════════════════
def build():
    prs = Presentation()
    prs.slide_width = _in(SW); prs.slide_height = _in(SH)

    # ─────────────────────────────── 1 · TITLE ──────────────────────────────
    s = slide(prs, NAVY)
    rect(s, 0, 0, SW, SH, fill=NAVY)
    # motif: concentric seismic arcs (rings) top-right
    for i, d in enumerate([7.5, 6.0, 4.6, 3.3]):
        oval(s, SW - d/2 - 1.2, -d/2 + 1.2, d, NAVY)  # base
        ring = s.shapes.add_shape(MSO_SHAPE.OVAL, _in(SW - d/2 - 1.2),
                                  _in(-d/2 + 1.2), _in(d), _in(d))
        ring.fill.background()
        ring.line.color.rgb = AMBER if i == 0 else (TEAL if i % 2 else NAVY2)
        ring.line.width = Pt(1.5 if i else 2.25)
        ring.shadow.inherit = False
    rect(s, 0, 0, 0.22, SH, fill=AMBER)
    text(s, 0.9, 1.05, 10, 0.4,
         [R("IŞIK UNIVERSITY  ·  FACULTY OF NATURAL SCIENCES AND ENGINEERING", 13, ICE, bold=True, font=BODY)])
    text(s, 0.9, 1.45, 10, 0.4,
         [R("Department of Industrial Engineering — INDE 4902 Graduation Design Project", 12.5, RGBColor(0x9F,0xB6,0xCC), font=BODY)])
    text(s, 0.88, 2.5, 11.2, 2.4,
         [[R("Design of a Spatial Decision", 40, WHITE, bold=True, font=HEAD)],
          [R("Support System for Earthquake", 40, WHITE, bold=True, font=HEAD)],
          [R("Assembly Area Planning", 40, AMBER, bold=True, font=HEAD)]],
         line=1.02, space_after=2)
    text(s, 0.9, 5.15, 11, 0.5,
         [R("A capacity-aware p-Median optimisation framework built on open spatial data, demonstrated on Istanbul's Kadıköy district.",
            15, ICE, italic=True, font=BODY)], line=1.1)
    rect(s, 0.9, 5.95, 5.6, 0.02, fill=RGBColor(0x2A,0x44,0x5E))
    text(s, 0.9, 6.15, 8, 0.8,
         [[R("Zeliha Karabay · Yusuf Kılıç · Sevcan Görgülü", 13.5, WHITE, bold=True, font=BODY)],
          [R("Supervisors: Seda Baş Güre & İsmail Kayahan   |   May 2026", 12, RGBColor(0x9F,0xB6,0xCC), font=BODY)]],
         space_after=3, line=1.1)

    # ─────────────────────────────── 2 · AGENDA ─────────────────────────────
    s = slide(prs); accent_side(s)
    kicker(s, 0.55, 0.5, "Presentation overview")
    title(s, 0.55, 0.82, 12, "Agenda")
    agenda = [
        ("Motivation & problem", "Istanbul's seismic risk and the planning gap"),
        ("Objectives & SDGs", "What we set out to build and why it matters"),
        ("System architecture", "A modular, reproducible spatial pipeline"),
        ("Data & classification", "From raw OSM tags to analysis-ready data"),
        ("Optimisation model", "Capacity-aware p-Median formulation"),
        ("Solver strategy", "Exact ILP and scalable greedy heuristic"),
        ("Case study & results", "Kadıköy: ~6,665 buildings, 154 areas"),
        ("Impact & conclusion", "Cost-benefit, SDGs, and future work"),
    ]
    x0, y0, cw, ch, gx, gy = 0.55, 1.95, 5.95, 1.12, 0.35, 0.25
    for i, (h, d) in enumerate(agenda):
        col, row = i % 2, i // 2
        x = x0 + col * (cw + gx); y = y0 + row * (ch + gy)
        card(s, x, y, cw, ch)
        chip(s, x + 0.32, y + ch/2 - 0.21, i + 1, color=AMBER if col == 0 else TEAL)
        text(s, x + 1.0, y + 0.2, cw - 1.2, ch - 0.3,
             [[R(h, 16, NAVY, bold=True, font=HEAD)],
              [R(d, 12.5, MUTED, font=BODY)]], space_after=2, line=1.02)
    footer(s, 2, TOTAL, "")

    # ─────────────────────────────── 3 · THE PROBLEM ────────────────────────
    s = slide(prs, NAVY)
    rect(s, 0, 0, SW, SH, fill=NAVY)
    rect(s, 0, 0, 0.22, SH, fill=AMBER)
    kicker(s, 0.9, 0.6, "Why this matters", color=AMBER)
    title(s, 0.9, 0.95, 11.5, "A megacity on a major fault line", color=WHITE)
    text(s, 0.9, 1.95, 11.4, 0.9,
         [R("Istanbul sits beside the North Anatolian Fault. In the first minutes after a major earthquake, residents must reach safe open spaces — assembly areas — on foot, often when roads and vehicles are unusable. How well these areas are placed directly affects how many lives are protected.",
            15.5, ICE, font=BODY)], line=1.18)
    stats = [
        ("15.6 M+", "residents to plan for"),
        ("39", "districts city-wide"),
        ("1.5 m²", "AFAD area per person"),
        ("On foot", "the only reliable post-quake mode"),
    ]
    x0 = 0.9; cw = 2.78; gx = 0.27; y = 3.4
    for i, (big, lab) in enumerate(stats):
        x = x0 + i * (cw + gx)
        card(s, x, y, cw, 1.95, fill=NAVY2, border=RGBColor(0x25,0x44,0x60), shadow=False)
        rect(s, x, y, cw, 0.12, fill=AMBER if i % 2 == 0 else TEAL, rounded=False)
        text(s, x + 0.25, y + 0.42, cw - 0.5, 0.9,
             [R(big, 33, WHITE, bold=True, font=HEAD)])
        text(s, x + 0.25, y + 1.25, cw - 0.5, 0.6,
             [R(lab, 12.5, ICE, font=BODY)], line=1.05)
    text(s, 0.9, 5.7, 11.4, 0.8,
         [R("Today, areas are often assigned by administrative habit — not by measuring whether residents can actually reach a large-enough area quickly. The true walking accessibility of the population is rarely quantified.",
            13.5, RGBColor(0x9F,0xB6,0xCC), italic=True, font=BODY)], line=1.15)
    footer(s, 3, TOTAL, "", dark=True)

    # ─────────────────────────────── 4 · TWO GAPS ───────────────────────────
    s = slide(prs); accent_side(s)
    kicker(s, 0.55, 0.5, "Problem description")
    title(s, 0.55, 0.82, 12, "Two gaps stand between data and decisions")
    gaps = [
        (AMBER, "1", "The data gap",
         "OpenStreetMap is an open “digital treasure”, but its crowd-sourced tags are raw and ambiguous.",
         ["Many structures are tagged only “building=yes” — no function.",
          "Manual cleaning + spatial joining can take days per district.",
          "Inconsistent tagging undermines any downstream model."]),
        (TEAL, "2", "The planning gap",
         "Assembly-area assignment is rarely backed by quantitative accessibility analysis.",
         ["Areas chosen by administrative boundary, not reachability.",
          "Capacity vs. real demand is seldom checked.",
          "Dense neighbourhoods may be served by distant, small areas."]),
    ]
    x0, y0, cw, ch = 0.55, 1.95, 5.95, 4.55
    for i, (c, n, h, sub, items) in enumerate(gaps):
        x = x0 + i * (cw + 0.35)
        card(s, x, y0, cw, ch)
        rect(s, x, y0, cw, 0.16, fill=c, rounded=False)
        chip(s, x + 0.4, y0 + 0.4, n, color=c)
        text(s, x + 1.05, y0 + 0.42, cw - 1.3, 0.6, [R(h, 21, NAVY, bold=True, font=HEAD)])
        text(s, x + 0.42, y0 + 1.25, cw - 0.84, 1.0,
             [R(sub, 14, INK, italic=True, font=BODY)], line=1.12)
        bullet_block(s, x + 0.42, y0 + 2.45, cw - 0.84, items, size=13.5,
                     marker_color=c, gap=10, lh=1.08)
    footer(s, 4, TOTAL, "")

    # ─────────────────────────────── 5 · OBJECTIVES + SDG ───────────────────
    s = slide(prs); accent_side(s)
    kicker(s, 0.55, 0.5, "Objectives & relevance")
    title(s, 0.55, 0.82, 12, "What we built — and why it matters")
    # left: objectives
    text(s, 0.55, 1.85, 6.4, 0.4, [R("Project objectives", 16, NAVY, bold=True, font=HEAD)])
    objs = [
        "Extract & clean OSM data; resolve ambiguous tags via a rule engine.",
        "Estimate building-level population as demand weight.",
        "Compute realistic pedestrian travel times on the street network.",
        "Solve a capacity-aware p-Median model (exact + heuristic).",
        "Deliver an interactive decision-support tool with exports.",
    ]
    tb = s.shapes.add_textbox(_in(0.55), _in(2.3), _in(6.3), _in(4.3))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0
    for i, o in enumerate(objs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(13); p.line_spacing = 1.06
        r0 = p.add_run(); r0.text = f"{i+1}   "
        r0.font.size = Pt(15); r0.font.bold = True; r0.font.name = HEAD; r0.font.color.rgb = AMBER
        r1 = p.add_run(); r1.text = o
        r1.font.size = Pt(14.5); r1.font.name = BODY; r1.font.color.rgb = INK
    # right: SDG cards
    text(s, 7.35, 1.85, 5.4, 0.4, [R("Aligned UN Sustainable Development Goals", 16, NAVY, bold=True, font=HEAD)])
    sdgs = [
        (RGBColor(0xFD,0x9D,0x24), "SDG 11", "Sustainable Cities & Communities",
         "Target 11.5 — reduce deaths and people affected by disasters."),
        (RGBColor(0xF3,0x6E,0x24), "SDG 9", "Industry, Innovation & Infrastructure",
         "A reusable, open-data analytical infrastructure."),
        (RGBColor(0x3F,0x7E,0x44), "SDG 13", "Climate Action",
         "Same accessibility logic supports climate-hazard evacuation."),
    ]
    yy = 2.35
    for c, code, name, desc in sdgs:
        card(s, 7.35, yy, 5.45, 1.28)
        rect(s, 7.35, yy, 0.16, 1.28, fill=c, rounded=False)
        text(s, 7.65, yy + 0.18, 1.5, 0.9,
             [R(code, 19, c, bold=True, font=HEAD)])
        text(s, 9.0, yy + 0.16, 3.7, 1.1,
             [[R(name, 13.5, NAVY, bold=True, font=BODY)],
              [R(desc, 11.5, MUTED, font=BODY)]], space_after=2, line=1.05)
        yy += 1.42
    footer(s, 5, TOTAL, "")

    # ─────────────────────────────── 6 · ARCHITECTURE (fig 4.1) ─────────────
    s = slide(prs); accent_side(s)
    kicker(s, 0.55, 0.5, "System design")
    title(s, 0.55, 0.82, 12, "A modular, reproducible architecture")
    # figure on left — height-bounded so it never collides with the caption
    img = FIG / "figure_4_1_sdss_architecture.png"
    fw = 6.1  # → height 6.1*(1480/2014)=4.48, image 2.02..6.50
    card(s, 0.5, 1.92, fw + 0.1, 4.62, fill=WHITE, border=HAIR, shadow=True)  # bottom 6.54
    s.shapes.add_picture(str(img), _in(0.6), _in(2.02), width=_in(fw))
    # right column: layer explanation
    layers = [
        (AMBER, "Configuration", "All constants and the tag/category rule registry in one place."),
        (TEAL, "Services", "OSM access, spatial processing, classification, export."),
        (NAVY, "Optimiser", "Data loader, OD matrix, population estimator, p-Median solvers."),
        (GREEN, "Presentation", "Two Streamlit apps + a consistent component design system."),
    ]
    x = 7.05; yy = 2.1
    for c, h, d in layers:
        oval(s, x, yy + 0.05, 0.34, c)
        text(s, x + 0.5, yy - 0.05, 5.4, 1.0,
             [[R(h, 15.5, NAVY, bold=True, font=HEAD)],
              [R(d, 12, MUTED, font=BODY)]], space_after=2, line=1.04)
        yy += 1.15
    text(s, 0.55, 6.68, 12, 0.35,
         [R("Figure 4.1 — High-level architecture of the spatial decision-support system. The optimisation core has no UI dependency, so it is independently testable and reusable.",
            10.5, MUTED, italic=True, font=BODY)], line=1.0)
    footer(s, 6, TOTAL, "")

    # ─────────────────────────────── 7 · METHODOLOGY PIPELINE ───────────────
    s = slide(prs); accent_side(s)
    kicker(s, 0.55, 0.5, "Methodology")
    title(s, 0.55, 0.82, 12, "An end-to-end pipeline in four stages")
    steps = [
        (AMBER, "Extract", "Buildings, candidate areas & boundaries from OSM via Overpass."),
        (TEAL, "Prepare", "Clean geometry, classify tags, estimate building population."),
        (NAVY, "Route", "Pedestrian OD travel-time matrix on the street network (OSMnx)."),
        (GREEN, "Optimise", "Capacity-aware p-Median → assignments, KPIs, maps, reports."),
    ]
    x0, y, cw, ch = 0.55, 1.95, 2.85, 2.25
    gap = (12.25 - 0.55 - 4*cw) / 3
    for i, (c, h, d) in enumerate(steps):
        x = x0 + i * (cw + gap)
        card(s, x, y, cw, ch)
        rect(s, x, y, cw, 0.7, fill=c, rounded=True)
        rect(s, x, y + 0.35, cw, 0.35, fill=c)  # square off bottom of header
        text(s, x, y + 0.12, cw, 0.5, [R(f"STAGE {i+1}", 12, WHITE, bold=True, font=BODY)],
             align=PP_ALIGN.CENTER)
        text(s, x + 0.25, y + 0.95, cw - 0.5, 0.5, [R(h, 19, NAVY, bold=True, font=HEAD)])
        text(s, x + 0.25, y + 1.5, cw - 0.5, 1.1, [R(d, 12.5, MUTED, font=BODY)], line=1.08)
        if i < 3:
            ar = s.shapes.add_shape(MSO_SHAPE.CHEVRON,
                                    _in(x + cw + gap/2 - 0.16), _in(y + ch/2 - 0.16),
                                    _in(0.32), _in(0.32))
            ar.fill.solid(); ar.fill.fore_color.rgb = AMBER
            ar.line.fill.background(); ar.shadow.inherit = False
    # figure 4.2 below — width-bound to keep height small (ratio 1180/2579)
    fw2 = 5.2   # → height 5.2*(1180/2579)=2.38, fits between cards and caption
    s.shapes.add_picture(str(FIG / "figure_4_2_optimization_workflow.png"),
                         _in((SW - fw2) / 2), _in(4.35), width=_in(fw2))
    text(s, 0.55, 6.92, 12, 0.3,
         [R("Figure 4.2 — Optimisation workflow from data loading to results.", 10.5, MUTED, italic=True, font=BODY)],
         align=PP_ALIGN.CENTER)
    footer(s, 7, TOTAL, "")

    # ─────────────────────────────── 8 · DATA & CLASSIFICATION ──────────────
    s = slide(prs); accent_side(s)
    kicker(s, 0.55, 0.5, "Stage 1 — Data & classification")
    title(s, 0.55, 0.82, 12, "Turning “building=yes” into meaning")
    # left text
    bullet_block(s, 0.55, 2.0, 6.0, [
        ("Rule-based engine — ", "resolves ambiguous tags using strict tags, supporting tags, and name patterns."),
        ("Multi-mirror Overpass — ", "queries retried across endpoints; partial/total failures reported, never silent."),
        ("Spatial enrichment — ", "each record deduplicated, geocoded to its neighbourhood, confidence-scored."),
        ("Boundary checks — ", "out-of-district leakage flagged and filterable."),
    ], size=14.5, gap=14, lh=1.1)
    # right: before/after compare
    card(s, 7.0, 1.95, 5.8, 1.85, fill=RGBColor(0xFB,0xEC,0xE6))
    rect(s, 7.0, 1.95, 0.16, 1.85, fill=AMBER)
    text(s, 7.3, 2.12, 5.3, 0.4, [R("RAW OSM TAG", 11.5, AMBER, bold=True, font=BODY)])
    text(s, 7.3, 2.5, 5.3, 0.5, [R("building = yes", 18, INK, bold=True, font="Consolas")])
    text(s, 7.3, 3.05, 5.3, 0.7,
         [R("Ambiguous — could be a home, school, hospital, or shop.", 12.5, MUTED, italic=True, font=BODY)], line=1.05)
    # arrow
    ar = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, _in(9.6), _in(3.92), _in(0.55), _in(0.42))
    ar.fill.solid(); ar.fill.fore_color.rgb = TEAL; ar.line.fill.background(); ar.shadow.inherit = False
    card(s, 7.0, 4.5, 5.8, 1.95, fill=RGBColor(0xE6,0xF2,0xF1))
    rect(s, 7.0, 4.5, 0.16, 1.95, fill=TEAL)
    text(s, 7.3, 4.66, 5.3, 0.4, [R("CLASSIFIED RECORD", 11.5, TEAL, bold=True, font=BODY)])
    text(s, 7.3, 5.04, 5.3, 0.5, [R("Healthcare · Hospital", 18, GREEN, bold=True, font="Consolas")])
    text(s, 7.3, 5.6, 5.3, 0.8,
         [R("Inferred from name “…Hastanesi” + supporting tags; tagged with confidence = High.",
            12.5, MUTED, italic=True, font=BODY)], line=1.05)
    footer(s, 8, TOTAL, "")

    # ─────────────────────────────── 9 · POPULATION + OD ────────────────────
    s = slide(prs); accent_side(s)
    kicker(s, 0.55, 0.5, "Stages 2–3 — Demand & travel time")
    title(s, 0.55, 0.82, 12, "Estimating demand and walking time")
    # two cards
    card(s, 0.55, 2.0, 5.95, 4.45)
    rect(s, 0.55, 2.0, 5.95, 0.16, fill=AMBER, rounded=False)
    oval(s, 0.9, 2.45, 0.5, AMBER); text(s, 0.9, 2.5, 0.5, 0.5, [R("◔", 17, WHITE, bold=True, font=BODY)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.6, 2.5, 4.6, 0.5, [R("Population estimation", 18, NAVY, bold=True, font=HEAD)])
    bullet_block(s, 0.9, 3.35, 5.3, [
        ("Footprint method — ", "area × floors × occupancy factor (0.025 person/m²)."),
        ("Uniform method — ", "distribute a TÜİK neighbourhood total across buildings."),
        ("Constants — ", "household size 3.24, default 4 floors when tag absent."),
    ], size=13.5, gap=11, lh=1.08)
    card(s, 6.85, 2.0, 5.95, 4.45)
    rect(s, 6.85, 2.0, 5.95, 0.16, fill=TEAL, rounded=False)
    oval(s, 7.2, 2.45, 0.5, TEAL); text(s, 7.2, 2.5, 0.5, 0.5, [R("◇", 17, WHITE, bold=True, font=BODY)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 7.9, 2.5, 4.6, 0.5, [R("Pedestrian OD matrix", 18, NAVY, bold=True, font=HEAD)])
    bullet_block(s, 7.2, 3.35, 5.3, [
        ("Real network — ", "walking graph from OSMnx, downloaded once and cached."),
        ("Shortest paths — ", "facility-to-building times at 4.8 km/h (level ground)."),
        ("Robust fallback — ", "haversine × 1.4 detour factor when the network is unavailable."),
    ], size=13.5, gap=11, lh=1.08)
    footer(s, 9, TOTAL, "")

    # ─────────────────────────────── 10 · MATH MODEL ────────────────────────
    s = slide(prs, NAVY)
    rect(s, 0, 0, SW, SH, fill=NAVY); rect(s, 0, 0, 0.22, SH, fill=AMBER)
    kicker(s, 0.9, 0.55, "The core — Stage 4", color=AMBER)
    title(s, 0.9, 0.9, 11.5, "Capacity-aware p-Median model", color=WHITE)
    # objective box
    card(s, 0.9, 1.95, 11.5, 1.15, fill=NAVY2, border=RGBColor(0x25,0x44,0x60))
    text(s, 1.2, 2.12, 11, 0.4, [R("OBJECTIVE  (efficiency, min-sum)", 12, AMBER, bold=True, font=BODY)])
    text(s, 1.2, 2.5, 11, 0.55,
         [R("minimise   Σᵢ Σⱼ  wᵢ · dᵢⱼ · xᵢⱼ", 20, WHITE, bold=True, font="Cambria")])
    # constraints grid
    cons = [
        ("Σⱼ xᵢⱼ = 1", "Each building served by exactly one area."),
        ("xᵢⱼ ≤ yⱼ", "Assign only to an opened area."),
        ("Σⱼ yⱼ = p", "Exactly p areas are opened."),
        ("Σᵢ wᵢ xᵢⱼ ≤ Cⱼ yⱼ", "Capacity respected (Cⱼ = area ÷ density)."),
    ]
    x0, y, cw, ch = 0.9, 3.3, 5.65, 1.0
    for i, (eq, d) in enumerate(cons):
        col, row = i % 2, i // 2
        x = x0 + col * (cw + 0.2); yy = y + row * (ch + 0.18)
        card(s, x, yy, cw, ch, fill=NAVY2, border=RGBColor(0x25,0x44,0x60), shadow=False)
        text(s, x + 0.25, yy + 0.12, cw - 0.5, 0.4, [R(eq, 16, ICE, bold=True, font="Cambria")])
        text(s, x + 0.25, yy + 0.55, cw - 0.5, 0.4, [R(d, 11.5, RGBColor(0x9F,0xB6,0xCC), font=BODY)])
    text(s, 0.9, 5.55, 11.5, 0.9,
         [R("Two alternative objectives expose the equity dimension: a min-max objective (worst single walk) and a population-weighted 95th-percentile objective — outlier-robust and recommended for decision support. Density defaults to 1.5 m²/person (AFAD), adjustable for sensitivity analysis.",
            13, ICE, italic=True, font=BODY)], line=1.18)
    footer(s, 10, TOTAL, "", dark=True)

    # ─────────────────────────────── 11 · SOLVER STRATEGY ───────────────────
    s = slide(prs); accent_side(s)
    kicker(s, 0.55, 0.5, "Solver strategy")
    title(s, 0.55, 0.82, 12, "Exact where affordable, scalable where needed")
    cols = [
        (NAVY, "Exact — ILP", "PuLP / CBC", [
            "Provably optimal under the constraints.",
            "Used at or below 5,000 buildings.",
            "Reports CBC status (Optimal / …).",
            "Optional faster back-ends (HiGHS).",
        ], "Preferred for small instances"),
        (AMBER, "Heuristic — Heuristic", "greedy + 1-swap", [
            "Greedy init + 1-swap local search.",
            "Capacity-aware; fast on large data.",
            "Reports converged vs. iteration-limit.",
            "~1–5% typical optimality gap.",
        ], "Used above the threshold"),
    ]
    x0, y, cw, ch = 0.55, 1.95, 5.95, 4.05
    for i, (c, h, sub, items, tag) in enumerate(cols):
        x = x0 + i * (cw + 0.35)
        card(s, x, y, cw, ch)
        rect(s, x, y, cw, 0.95, fill=c, rounded=True); rect(s, x, y + 0.5, cw, 0.45, fill=c)
        text(s, x + 0.4, y + 0.16, cw - 0.8, 0.5, [R(h, 20, WHITE, bold=True, font=HEAD)])
        text(s, x + 0.4, y + 0.58, cw - 0.8, 0.35, [R(sub, 12.5, RGBColor(0xE7,0xEE,0xF6), italic=True, font="Consolas")])
        bullet_block(s, x + 0.42, y + 1.2, cw - 0.84, items, size=13.5,
                     marker_color=c, gap=10, lh=1.06)
        rect(s, x + 0.0, y + ch + 0.18, cw, 0.5, fill=RGBColor(0xEE,0xF3,0xF8), rounded=True)
        text(s, x, y + ch + 0.26, cw, 0.4, [R(tag, 12.5, c, bold=True, font=BODY)], align=PP_ALIGN.CENTER)
    # center "auto" badge
    oval(s, SW/2 - 0.55, y + 1.55, 1.1, WHITE, line=AMBER)
    text(s, SW/2 - 0.55, y + 1.7, 1.1, 0.8,
         [[R("AUTO", 13, NAVY, bold=True, font=HEAD)],
          [R("by size", 10.5, MUTED, font=BODY)]], align=PP_ALIGN.CENTER, space_after=0, line=0.95)
    footer(s, 11, TOTAL, "")

    # ─────────────────────────────── 12 · THE APPLICATIONS (UI) ─────────────
    s = slide(prs); accent_side(s)
    kicker(s, 0.55, 0.5, "Delivery")
    title(s, 0.55, 0.82, 12, "Two interactive decision-support apps")
    placeholder(s, 0.55, 2.0, 6.0, 3.05, "Screenshot — Data Extraction page (district + categories, live log, KPIs)")
    placeholder(s, 6.8, 2.0, 6.0, 3.05, "Screenshot — Optimisation tool (4-step workflow, map + coverage KPIs)")
    feats = [
        "Targeted OSM extraction across 39 districts & 100+ subcategories",
        "Interactive maps, analytics dashboard, name-lookup & boundary sync",
        "Multi-sheet Excel / CSV / GeoJSON exports — decision-ready outputs",
    ]
    bullet_block(s, 0.55, 5.4, 12.2, feats, size=13.5, gap=7, lh=1.05)
    footer(s, 12, TOTAL, "")

    # ─────────────────────────────── 13 · CASE STUDY KADIKOY ────────────────
    s = slide(prs, NAVY)
    rect(s, 0, 0, SW, SH, fill=NAVY); rect(s, 0, 0, 0.22, SH, fill=AMBER)
    kicker(s, 0.9, 0.55, "Case study", color=AMBER)
    title(s, 0.9, 0.9, 11.5, "Kadıköy — a representative district", color=WHITE)
    text(s, 0.9, 1.9, 11.4, 0.7,
         [R("The methodology is demonstrated in depth on Kadıköy, while the system is engineered as a general infrastructure that runs on any of Istanbul's 39 districts without modification.",
            14.5, ICE, font=BODY)], line=1.15)
    kp = [("≈ 6,665", "buildings (demand)"), ("154", "candidate assembly areas"),
          ("≈ 244,000", "estimated residents"), ("> threshold", "ideal ILP-vs-heuristic test")]
    x0 = 0.9; cw = 2.78; gx = 0.27; y = 3.0
    for i, (b, l) in enumerate(kp):
        x = x0 + i * (cw + gx)
        card(s, x, y, cw, 1.85, fill=NAVY2, border=RGBColor(0x25,0x44,0x60), shadow=False)
        rect(s, x, y, cw, 0.12, fill=TEAL if i % 2 else AMBER)
        text(s, x + 0.2, y + 0.42, cw - 0.4, 0.7, [R(b, 28, WHITE, bold=True, font=HEAD)])
        text(s, x + 0.2, y + 1.18, cw - 0.4, 0.55, [R(l, 12, ICE, font=BODY)], line=1.0)
    placeholder_dark = True
    rect(s, 0.9, 5.2, 11.5, 1.35, fill=NAVY2, line=TEAL, line_w=1.25, rounded=True)
    text(s, 0.9, 5.5, 11.5, 0.4, [R("🗺  MAP PLACEHOLDER", 12.5, TEAL, bold=True, font=BODY)], align=PP_ALIGN.CENTER)
    text(s, 0.9, 5.9, 11.5, 0.5,
         [R("Assignment map of Kadıköy — buildings coloured by their assigned assembly area (to be exported from the tool).",
            12, ICE, italic=True, font=BODY)], align=PP_ALIGN.CENTER, line=1.0)
    footer(s, 13, TOTAL, "", dark=True)

    # ─────────────────────────────── 14 · RESULTS ───────────────────────────
    s = slide(prs); accent_side(s)
    kicker(s, 0.55, 0.5, "Computational results")
    title(s, 0.55, 0.82, 12, "Coverage, capacity & solver trade-offs")
    placeholder(s, 0.55, 2.0, 6.0, 4.4, "Chart — population coverage (%) within 5 / 10 min vs. number of opened areas p")
    # right: result table scaffold
    text(s, 6.85, 1.95, 5.9, 0.4, [R("Base optimisation — Kadıköy (to be filled)", 14.5, NAVY, bold=True, font=HEAD)])
    rows = [
        ("Indicator", "Value", True),
        ("Opened areas (p)", "—", False),
        ("Pop-weighted avg. time", "— min", False),
        ("Population-weighted p95", "— min", False),
        ("Population ≤ 5 min", "— %", False),
        ("Population ≤ 10 min", "— %", False),
        ("Solver / status", "—", False),
        ("Runtime", "— s", False),
    ]
    ty, rh = 2.45, 0.46
    for i, (a, b, head) in enumerate(rows):
        yy = ty + i * rh
        fill = NAVY if head else (RGBColor(0xEE,0xF3,0xF8) if i % 2 else WHITE)
        rect(s, 6.85, yy, 5.9, rh, fill=fill, line=HAIR, line_w=0.75)
        tcol = WHITE if head else INK
        text(s, 7.05, yy + 0.08, 3.7, rh, [R(a, 12.5, tcol, bold=head, font=BODY)], anchor=MSO_ANCHOR.MIDDLE)
        text(s, 10.7, yy + 0.08, 1.9, rh, [R(b, 12.5, (WHITE if head else AMBER), bold=True, font=BODY)],
             anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
    text(s, 0.55, 6.55, 12.2, 0.4,
         [R("Note — sensitivity to p, capacity ON/OFF, density sweep, and ILP-vs-Heuristic tables are produced by running the Kadıköy experiments; figures to be inserted here.",
            10.5, MUTED, italic=True, font=BODY)], line=1.0)
    footer(s, 14, TOTAL, "")

    # ─────────────────────────────── 15 · TRANSPARENCY / RIGOR ──────────────
    s = slide(prs); accent_side(s)
    kicker(s, 0.55, 0.5, "Engineering rigor")
    title(s, 0.55, 0.82, 12, "Built to be trusted, not just to run")
    items = [
        (AMBER, "Honest solver reporting", "Optimal vs. heuristic, convergence vs. iteration-limit, and the reason for any fallback are all surfaced."),
        (TEAL, "Stale-result guard", "The UI warns when displayed KPIs no longer match changed inputs."),
        (NAVY, "Security by default", "User- and OSM-derived text is escaped before rendering in maps and reports."),
        (GREEN, "Reproducible & tested", "507 passing regression tests, CI on Python 3.10–3.12, pinned lock files."),
    ]
    x0, y0, cw, ch = 0.55, 1.95, 5.95, 2.15
    for i, (c, h, d) in enumerate(items):
        col, row = i % 2, i // 2
        x = x0 + col * (cw + 0.35); y = y0 + row * (ch + 0.25)
        card(s, x, y, cw, ch)
        oval(s, x + 0.35, y + 0.35, 0.55, c)
        text(s, x + 0.35, y + 0.41, 0.55, 0.55, [R("✓", 18, WHITE, bold=True, font=BODY)],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 1.15, y + 0.38, cw - 1.45, 0.5, [R(h, 16.5, NAVY, bold=True, font=HEAD)])
        text(s, x + 1.15, y + 0.92, cw - 1.45, 1.1, [R(d, 13, MUTED, font=BODY)], line=1.1)
    footer(s, 15, TOTAL, "")

    # ─────────────────────────────── 16 · COST-BENEFIT & IMPACT ─────────────
    s = slide(prs); accent_side(s)
    kicker(s, 0.55, 0.5, "Impact")
    title(s, 0.55, 0.82, 12, "Low cost, high value, broad reach")
    # left cost-benefit
    card(s, 0.55, 2.0, 5.95, 4.45)
    rect(s, 0.55, 2.0, 5.95, 0.16, fill=AMBER)
    text(s, 0.9, 2.4, 5.3, 0.4, [R("Cost–benefit", 18, NAVY, bold=True, font=HEAD)])
    bullet_block(s, 0.9, 3.0, 5.3, [
        ("Near-zero inputs — ", "open data (OSM) + open-source software; no licences."),
        ("Reusable — ", "generalises to all 39 districts at near-zero marginal cost."),
        ("Benefit — ", "shorter, capacity-feasible evacuation walks → reduced risk to life."),
    ], size=13.5, gap=13, lh=1.1)
    # right impact pillars
    card(s, 6.85, 2.0, 5.95, 4.45)
    rect(s, 6.85, 2.0, 5.95, 0.16, fill=TEAL)
    text(s, 7.2, 2.4, 5.3, 0.4, [R("Economic · Social · Ethical · Environmental", 15.5, NAVY, bold=True, font=HEAD)])
    bullet_block(s, 7.2, 3.0, 5.3, [
        ("Social — ", "quantifies and helps correct spatial inequities in provision."),
        ("Ethical — ", "transparent assumptions; optimal vs. heuristic never overstated."),
        ("Environmental — ", "uses existing open data & light compute; protects open/green space."),
    ], size=13.5, gap=13, lh=1.1)
    footer(s, 16, TOTAL, "")

    # ─────────────────────────────── 17 · LIMITATIONS & FUTURE ──────────────
    s = slide(prs); accent_side(s)
    kicker(s, 0.55, 0.5, "Honest boundaries")
    title(s, 0.55, 0.82, 12, "Limitations & future work")
    card(s, 0.55, 2.0, 5.95, 4.45, fill=RGBColor(0xFB,0xEC,0xE6))
    rect(s, 0.55, 2.0, 0.16, 4.45, fill=AMBER)
    text(s, 0.95, 2.35, 5.4, 0.4, [R("Limitations", 18, NAVY, bold=True, font=HEAD)])
    bullet_block(s, 0.95, 3.0, 5.2, [
        "OSM attributes are volunteer-contributed — may be missing.",
        "Population is estimated, not measured per resident.",
        "Constant walking speed on level ground (no slope/stairs).",
        "Candidate areas inferred from tags, not an official register.",
        "Deterministic — no demand uncertainty or network damage yet.",
    ], size=13.5, gap=11, lh=1.06, marker_color=AMBER)
    card(s, 6.85, 2.0, 5.95, 4.45, fill=RGBColor(0xE6,0xF2,0xF1))
    rect(s, 6.85, 2.0, 0.16, 4.45, fill=TEAL)
    text(s, 7.25, 2.35, 5.4, 0.4, [R("Future work", 18, NAVY, bold=True, font=HEAD)])
    bullet_block(s, 7.25, 3.0, 5.2, [
        "Validate against official AFAD assembly-area designations.",
        "Scale to the largest districts and a full city-wide run.",
        "Add finer demographics + scenario-based network damage.",
        "Unify the two apps into one multi-page interface.",
        "Add a bilingual (Turkish / English) presentation layer.",
    ], size=13.5, gap=11, lh=1.06, marker_color=TEAL)
    footer(s, 17, TOTAL, "")

    # ─────────────────────────────── 18 · CONCLUSION ────────────────────────
    s = slide(prs, NAVY)
    rect(s, 0, 0, SW, SH, fill=NAVY); rect(s, 0, 0, 0.22, SH, fill=AMBER)
    for i, d in enumerate([5.5, 4.2, 3.0]):
        ring = s.shapes.add_shape(MSO_SHAPE.OVAL, _in(-d/2 + 1.0), _in(SH - d/2 - 0.6), _in(d), _in(d))
        ring.fill.background(); ring.line.color.rgb = (AMBER if i == 0 else TEAL)
        ring.line.width = Pt(1.5); ring.shadow.inherit = False
    kicker(s, 0.9, 0.7, "Conclusion", color=AMBER)
    title(s, 0.9, 1.05, 11.5, "From open data to defensible decisions", color=WHITE, size=34)
    text(s, 0.9, 2.3, 11.3, 1.6,
         [R("We designed and built a reproducible spatial decision-support system that integrates open-data extraction and rule-based classification, building-level population estimation, network-based pedestrian routing, and a capacity-aware p-Median optimisation with both an exact and a scalable solver — delivered as interactive applications with maps, KPIs, and exportable reports.",
            16, ICE, font=BODY)], line=1.22)
    takeaways = [
        ("Open", "Only OSM data + open-source software."),
        ("Transparent", "Every assumption & solution quality is reported."),
        ("Scalable", "Demonstrated on Kadıköy; ready for all 39 districts."),
    ]
    x0 = 0.9; cw = 3.7; gx = 0.3; y = 4.6
    for i, (h, d) in enumerate(takeaways):
        x = x0 + i * (cw + gx)
        card(s, x, y, cw, 1.7, fill=NAVY2, border=RGBColor(0x25,0x44,0x60), shadow=False)
        rect(s, x, y, cw, 0.12, fill=AMBER if i == 0 else TEAL)
        text(s, x + 0.3, y + 0.35, cw - 0.6, 0.5, [R(h, 19, WHITE, bold=True, font=HEAD)])
        text(s, x + 0.3, y + 0.92, cw - 0.6, 0.7, [R(d, 12.5, ICE, font=BODY)], line=1.05)
    footer(s, 18, TOTAL, "", dark=True)

    # ─────────────────────────────── 19 · THANK YOU ─────────────────────────
    s = slide(prs, NAVY)
    rect(s, 0, 0, SW, SH, fill=NAVY)
    rect(s, 0, 0, 0.22, SH, fill=AMBER)
    for i, d in enumerate([8.5, 6.8, 5.1, 3.4]):
        ring = s.shapes.add_shape(MSO_SHAPE.OVAL, _in(SW - d/2 - 1.0), _in(SH/2 - d/2), _in(d), _in(d))
        ring.fill.background()
        ring.line.color.rgb = AMBER if i == 0 else (TEAL if i % 2 else RGBColor(0x25,0x44,0x60))
        ring.line.width = Pt(1.5); ring.shadow.inherit = False
    text(s, 0.9, 2.7, 11, 1.2, [R("Thank you", 52, WHITE, bold=True, font=HEAD)])
    text(s, 0.92, 3.95, 11, 0.5,
         [R("Questions & discussion welcome.", 18, AMBER, italic=True, font=BODY)])
    rect(s, 0.95, 4.75, 5.0, 0.02, fill=RGBColor(0x2A,0x44,0x5E))
    text(s, 0.92, 4.95, 11.5, 1.2,
         [[R("Zeliha Karabay · Yusuf Kılıç · Sevcan Görgülü", 14, WHITE, bold=True, font=BODY)],
          [R("Supervisors: Seda Baş Güre & İsmail Kayahan", 12.5, ICE, font=BODY)],
          [R("Işık University · Department of Industrial Engineering · INDE 4902 · May 2026", 11.5, RGBColor(0x9F,0xB6,0xCC), font=BODY)]],
         space_after=4, line=1.1)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"WROTE {OUT} ({OUT.stat().st_size:,} bytes, {len(prs.slides.__iter__.__self__._sldIdLst)} slides)")


if __name__ == "__main__":
    build()
